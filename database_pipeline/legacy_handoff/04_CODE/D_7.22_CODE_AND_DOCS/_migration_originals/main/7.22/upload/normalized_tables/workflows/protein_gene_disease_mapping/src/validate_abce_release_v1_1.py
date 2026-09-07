from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from .common import load_config
from .compare_evidence_filter_rules import RELATION_KEY, sha256


EXPECTED_SHEETS = [
    "1_Rule_B_Direct",
    "2_Primary_Disease",
    "3_Top10_Diseases",
    "4_Evidence_Level_Summary",
    "5_Source_Family_Summary",
    "6_Protein_Gene_Mapping",
    "7_Disease_Summary",
    "8_No_Rule_B_Disease",
    "9_QC_Report",
    "10_Metadata",
]


def require(condition: bool, message: str) -> None:
    if not condition:
        raise RuntimeError(message)


def key_set(frame: pd.DataFrame) -> set[tuple[str, str, str]]:
    return set(frame[RELATION_KEY].astype(str).itertuples(index=False, name=None))


def validate_checksums(root: Path, checksum_file: Path) -> int:
    lines = [line for line in checksum_file.read_text(encoding="utf-8").splitlines() if line.strip()]
    for line in lines:
        expected, relative = line.split("  ", 1)
        path = root / relative
        require(path.exists(), f"Checksum target is missing: {relative}")
        require(sha256(path) == expected, f"Checksum mismatch: {relative}")
    return len(lines)


def validate_release(config_path: str, scope: str = "archive") -> dict[str, Any]:
    root, config = load_config(config_path)
    release = root / config["outputs"]["release_dir"]
    b_parquet = release / "protein_gene_disease_rule_b_scored_v1.parquet"
    b_gzip = release / "protein_gene_disease_rule_b_scored_v1.tsv.gz"
    primary_path = release / "protein_primary_disease_rule_b_v1.tsv"
    top10_path = release / "protein_gene_disease_rule_b_top10_v1.tsv"
    workbook_path = release / "protein_gene_disease_rule_b_summary_v1.xlsx"
    for path in [b_parquet, b_gzip, primary_path, top10_path, workbook_path]:
        require(path.exists(), f"Required Rule B file is missing: {path.name}")

    rule_b = pd.read_parquet(b_parquet)
    rule_b_tsv = pd.read_csv(b_gzip, sep="\t", dtype=str, keep_default_na=False)
    require(len(rule_b) == 6978, f"Rule B relationship count is {len(rule_b)}, expected 6978")
    require(rule_b["target_uniprot_id"].nunique() == 1119, "Rule B protein count is not 1119")
    require(rule_b["disease_id"].nunique() == 2310, "Rule B disease count is not 2310")
    require(rule_b.duplicated(RELATION_KEY).sum() == 0, "Rule B has duplicate relation keys")
    require(key_set(rule_b) == key_set(rule_b_tsv), "Rule B Parquet and TSV.GZ keys differ")
    require(set(rule_b["evidence_level"]) == {"medium", "high", "very_high"}, "Rule B levels are invalid")
    require(not rule_b["evidence_level"].isin({"low", "not_included"}).any(), "Rule B contains excluded levels")
    require(rule_b["rule_b_pass"].all(), "Rule B contains a row without rule_b_pass")
    require((rule_b["rule_e_pass"] <= rule_b["rule_c_pass"]).all(), "Rule E is not a subset of Rule C")
    require((rule_b["rule_c_pass"] <= rule_b["rule_b_pass"]).all(), "Rule C is not a subset of Rule B")
    require(not rule_b["unknown_source_flag"].any(), "Unknown-source row passed Rule B")
    require(not rule_b["weak_sources_only_flag"].any(), "Weak-source-only row passed Rule B")
    require(rule_b["entity_type"].eq("disease").all(), "Non-disease row passed Rule B")

    primary = pd.read_csv(primary_path, sep="\t", dtype=str, keep_default_na=False)
    top10 = pd.read_csv(top10_path, sep="\t", dtype=str, keep_default_na=False)
    require(primary["target_uniprot_id"].is_unique, "Primary summary has duplicate proteins")
    no_disease = primary["primary_status"].eq("no_rule_b_supported_disease")
    require(primary.loc[no_disease, "primary_disease_id"].eq("").all(), "No-Rule-B protein received a fallback disease")
    require(top10.groupby("target_uniprot_id").size().max() <= 10, "Top-10 summary exceeds ten rows per protein")
    require(set(top10["evidence_level"]).issubset({"medium", "high", "very_high"}), "Top-10 contains a low row")
    require(set(top10["disease_id"]).issubset(set(rule_b["disease_id"])), "Top-10 contains a disease outside Rule B")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    require(workbook.sheetnames == EXPECTED_SHEETS, f"Unexpected workbook sheets: {workbook.sheetnames}")
    workbook.close()
    require(workbook_path.stat().st_size < 100 * 1024 * 1024, "Rule B workbook exceeds 100 MiB")

    result: dict[str, Any] = {
        "status": "passed",
        "scope": scope,
        "rule_b_relationship_count": len(rule_b),
        "rule_b_protein_count": int(rule_b["target_uniprot_id"].nunique()),
        "rule_b_disease_count": int(rule_b["disease_id"].nunique()),
        "rule_b_duplicate_key_count": int(rule_b.duplicated(RELATION_KEY).sum()),
        "level_counts": {key: int(value) for key, value in rule_b["evidence_level"].value_counts().sort_index().items()},
        "primary_protein_count": len(primary),
        "primary_without_rule_b_count": int(no_disease.sum()),
        "top10_row_count": len(top10),
        "workbook_bytes": workbook_path.stat().st_size,
        "rule_b_checksum_entry_count": validate_checksums(root, release / "rule_b_checksums.sha256"),
    }

    if scope == "archive":
        scored_path = release / "protein_gene_disease_abce_scored_v1.parquet"
        require(scored_path.exists(), "Archive scored Parquet is missing")
        scored = pd.read_parquet(scored_path)
        counts = {rule: int(scored[f"rule_{rule.lower()}_pass"].sum()) for rule in "ABCE"}
        require(counts == {"A": 63269, "B": 6978, "C": 3860, "E": 2384}, f"Archive rule counts differ: {counts}")
        require((scored["rule_b_pass"] <= scored["rule_a_pass"]).all(), "Rule B is not a subset of Rule A")
        require((scored["rule_c_pass"] <= scored["rule_b_pass"]).all(), "Rule C is not a subset of Rule B")
        require((scored["rule_e_pass"] <= scored["rule_c_pass"]).all(), "Rule E is not a subset of Rule C")
        require(set(scored["evidence_level"]) == {"low", "medium", "high", "very_high"}, "Archive levels are incomplete")
        manifest = json.loads((release / "ABCE_SCORING_MANIFEST.json").read_text(encoding="utf-8"))
        require(manifest["historical_rule_d_used_for_scoring"] is False, "Historical Rule D is marked as formal")
        result["rule_counts"] = counts
        result["archive_level_counts"] = {
            key: int(value) for key, value in scored["evidence_level"].value_counts().sort_index().items()
        }
        result["comparison_material_count"] = len(manifest["comparison_materials"])
    return result
