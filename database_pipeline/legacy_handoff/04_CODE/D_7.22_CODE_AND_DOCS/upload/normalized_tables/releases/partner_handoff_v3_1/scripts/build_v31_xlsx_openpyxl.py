import csv
import hashlib
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(r"C:\github-repos\upload\normalized_tables\outputs\v3_1_completion")
OUTPUT = BASE / "normalized_tables_v3_1.xlsx"
csv.field_size_limit(100_000_000)

SHEETS = [
    ("1_Binary_Relationships", "drug_protein_binary_relationships_v3_1.tsv", 101802),
    ("2_Protein_Database", "protein_database_v3_1.tsv", 1773),
    ("3_Small_Molecules", "small_molecule_database_v3_1.tsv", 87611),
    ("4_Pocket_Instances", "pocket_instances_v3_1.tsv", 6815),
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_ALIGNMENT = Alignment(wrap_text=True, vertical="center")


def sha256(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def safe_cell(ws, value):
    if isinstance(value, str) and value.startswith("="):
        cell = WriteOnlyCell(ws, value=value)
        cell.data_type = "s"
        return cell
    return value


def styled_header(ws, values):
    out = []
    for value in values:
        cell = WriteOnlyCell(ws, value=value)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT
        out.append(cell)
    return out


def build_sheet(wb, sheet_name, file_name, expected_rows):
    path = BASE / file_name
    ws = wb.create_sheet(sheet_name)
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = True

    with path.open(newline="", encoding="utf-8-sig") as f:
        reader = csv.reader(f, delimiter="\t")
        header = next(reader)
        sample = []
        widths = [len(str(x)) for x in header]
        for _ in range(249):
            try:
                row = next(reader)
            except StopIteration:
                break
            sample.append(row)
            for i, value in enumerate(row):
                if i < len(widths):
                    widths[i] = max(widths[i], min(len(str(value)), 40))

        last_col = get_column_letter(len(header))
        ws.auto_filter.ref = f"A1:{last_col}{expected_rows + 1}"
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = max(10, min(width + 2, 42))
        ws.row_dimensions[1].height = 30

        ws.append(styled_header(ws, header))
        written = 0
        for row in sample:
            ws.append([safe_cell(ws, x) for x in row])
            written += 1
        for row in reader:
            ws.append([safe_cell(ws, x) for x in row])
            written += 1

    if written != expected_rows:
        raise RuntimeError(f"{sheet_name}: expected {expected_rows} rows, wrote {written}")
    print(json.dumps({"sheet": sheet_name, "rows": written, "cols": len(header)}, ensure_ascii=False), flush=True)


def main():
    wb = Workbook(write_only=True)
    for sheet_name, file_name, expected_rows in SHEETS:
        build_sheet(wb, sheet_name, file_name, expected_rows)
    wb.save(OUTPUT)

    check = load_workbook(OUTPUT, read_only=True, data_only=False)
    validation = {
        "output": str(OUTPUT),
        "bytes": OUTPUT.stat().st_size,
        "sha256": sha256(OUTPUT),
        "sheets": [],
        "formula_cells_in_sampled_rows": 0,
    }
    for sheet_name, _, expected_rows in SHEETS:
        ws = check[sheet_name]
        formula_count = 0
        for row in ws.iter_rows(min_row=1, max_row=1000, values_only=False):
            for cell in row:
                if cell.data_type == "f":
                    formula_count += 1
        actual_total_rows = sum(1 for _ in ws.iter_rows(values_only=True))
        header_values = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        validation["formula_cells_in_sampled_rows"] += formula_count
        validation["sheets"].append({
            "name": sheet_name,
            "data_rows": actual_total_rows - 1,
            "cols": len(header_values),
            "freeze_panes": str(ws.freeze_panes),
            "auto_filter": ws.auto_filter.ref,
            "header_fill": ws[1][0].fill.fgColor.rgb,
            "header_font_color": ws[1][0].font.color.rgb if ws[1][0].font.color and ws[1][0].font.color.type == "rgb" else "",
            "expected_rows": expected_rows,
        })
    check.close()
    validation["validation_passed"] = (
        validation["formula_cells_in_sampled_rows"] == 0
        and all(x["data_rows"] == x["expected_rows"] for x in validation["sheets"])
        and all(x["freeze_panes"] == "A2" for x in validation["sheets"])
        and all(bool(x["auto_filter"]) for x in validation["sheets"])
    )
    (BASE / "V3_1_XLSX_VALIDATION_REPORT.json").write_text(json.dumps(validation, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(validation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
