from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

import pdfplumber
from docx import Document
from docx.oxml.ns import qn


ROOT = Path(r"C:\Users\Administrator\HuMemLigDB_manuscript")
DOCX = ROOT / "HuMemLigDB_NAR_manuscript_draft_v0.1.docx"
PDF = ROOT / "rendered_v0_1" / "HuMemLigDB_NAR_manuscript_draft_v0.1_qa2.pdf"
REPORT = ROOT / "HuMemLigDB_NAR_manuscript_draft_v0.1_QA.json"


def compact(text: str, limit: int = 180) -> str:
    return re.sub(r"\s+", " ", text).strip()[:limit]


def inspect_pdf() -> dict:
    page_reports = []
    all_text = []
    global_errors = []
    with pdfplumber.open(PDF) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            words = page.extract_words(extra_attrs=["fontname", "size"], keep_blank_chars=False)
            text = page.extract_text(x_tolerance=2, y_tolerance=2) or ""
            all_text.append(text)
            lines = [compact(x) for x in text.splitlines() if x.strip()]
            outside = []
            for word in words:
                x0 = float(word["x0"])
                x1 = float(word["x1"])
                top = float(word["top"])
                bottom = float(word["bottom"])
                if x0 < -0.5 or top < -0.5 or x1 > page.width + 0.5 or bottom > page.height + 0.5:
                    outside.append([x0, top, x1, bottom, word["text"]])
            sizes = [round(float(w.get("size") or 0), 2) for w in words if float(w.get("size") or 0) > 0]
            fonts = sorted({str(w.get("fontname") or "") for w in words if w.get("fontname")})
            errors = []
            if not words:
                errors.append("blank_page")
            if outside:
                errors.append("text_outside_page")
            if sizes and min(sizes) < 7.5:
                errors.append("font_below_7_5_pt")
            # Detect text pressed implausibly close to non-header/footer page edges.
            body_words = [w for w in words if 35 <= float(w["top"]) <= page.height - 35]
            if any(float(w["x0"]) < 45 or float(w["x1"]) > page.width - 45 for w in body_words):
                errors.append("body_text_near_lateral_edge")
            page_reports.append(
                {
                    "page": page_number,
                    "size_points": [page.width, page.height],
                    "word_count": len(words),
                    "min_font_pt": min(sizes) if sizes else None,
                    "max_font_pt": max(sizes) if sizes else None,
                    "fonts": fonts,
                    "first_lines": lines[:5],
                    "last_lines": lines[-5:],
                    "outside_words": outside[:10],
                    "errors": sorted(set(errors)),
                }
            )
            global_errors.extend(f"page_{page_number}:{x}" for x in sorted(set(errors)))
    full_text = "\n".join(all_text)
    prohibited = [x for x in ("294,?", "[This source-count", "**", "```", "PLACEHOLDER") if x in full_text]
    required = [
        "Abstract",
        "Introduction",
        "Results",
        "Discussion",
        "Materials and methods",
        "Data availability",
        "1,305,791",
        "7,904",
        "3,003,306",
    ]
    return {
        "page_count": len(page_reports),
        "prohibited_tokens_found": prohibited,
        "required_tokens_missing": [x for x in required if x not in full_text],
        "global_errors": global_errors,
        "pages": page_reports,
    }


def inspect_docx() -> dict:
    doc = Document(DOCX)
    section = doc.sections[0]
    text = "\n".join(p.text for p in doc.paragraphs)
    errors = []
    table_reports = []
    for idx, table in enumerate(doc.tables, start=1):
        widths = [int(x.get(qn("w:w"))) for x in table._tbl.tblGrid.findall(qn("w:gridCol"))]
        tbl_w = table._tbl.tblPr.find(qn("w:tblW"))
        tbl_ind = table._tbl.tblPr.find(qn("w:tblInd"))
        width = int(tbl_w.get(qn("w:w"))) if tbl_w is not None else None
        indent = int(tbl_ind.get(qn("w:w"))) if tbl_ind is not None else None
        if sum(widths) != 9360 or width != 9360:
            errors.append(f"table_{idx}:geometry")
        if indent != 120:
            errors.append(f"table_{idx}:indent")
        table_reports.append(
            {
                "table": idx,
                "rows": len(table.rows),
                "cols": len(table.columns),
                "grid_widths": widths,
                "grid_sum": sum(widths),
                "tblW": width,
                "tblInd": indent,
            }
        )
    if "294,?" in text or "[This source-count" in text:
        errors.append("unresolved_source_count_placeholder")
    if "1,305,791" not in text:
        errors.append("core_pair_count_missing")
    with zipfile.ZipFile(DOCX) as archive:
        corrupt = archive.testzip()
        if corrupt:
            errors.append(f"zip_crc:{corrupt}")
    return {
        "page_size_inches": [section.page_width.inches, section.page_height.inches],
        "margins_inches": [
            section.top_margin.inches,
            section.right_margin.inches,
            section.bottom_margin.inches,
            section.left_margin.inches,
        ],
        "paragraph_count": len(doc.paragraphs),
        "tables": table_reports,
        "errors": errors,
    }


def main() -> None:
    result = {"docx": inspect_docx(), "pdf": inspect_pdf()}
    hard_pdf_errors = [x for x in result["pdf"]["global_errors"] if "blank_page" in x or "text_outside_page" in x]
    result["status"] = (
        "PASS"
        if not result["docx"]["errors"]
        and not hard_pdf_errors
        and not result["pdf"]["prohibited_tokens_found"]
        and not result["pdf"]["required_tokens_missing"]
        else "REVIEW"
    )
    REPORT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(REPORT)
    print(result["status"])
    print("\n".join(result["pdf"]["global_errors"]))


if __name__ == "__main__":
    main()
