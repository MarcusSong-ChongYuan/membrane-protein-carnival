from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


BASE = Path(r"C:\github-repos\upload\normalized_tables")
SRC = BASE / "data_dictionary_review.tsv"
OUT_TSV = BASE / "data_dictionary_review_zh.tsv"
OUT_XLSX = BASE / "data_dictionary_review_zh.xlsx"


HEADERS_ZH = [
    "所属Sheet",
    "对应TSV文件",
    "列序号",
    "字段名",
    "中文字段名",
    "字段说明",
    "数据类型",
    "是否主键",
    "是否外键",
    "外键指向",
    "允许值",
    "多值分隔符",
    "缺失值含义",
    "字段角色",
    "来源或生成规则",
    "总行数",
    "真实有值数",
    "not_found计数",
    "空值数",
    "真实覆盖率%",
    "观察到的唯一值数",
    "高频取值",
    "示例值",
    "备注",
    "推荐用法",
]


TYPE_ZH = {
    "identifier": "标识符",
    "multi_identifier": "多值标识符",
    "text": "文本",
    "integer": "整数",
    "float": "数值",
    "boolean": "布尔值",
    "category": "分类值",
    "multi_category": "多值分类",
}


ROLE_ZH = {
    "core_identifier": "核心标识字段",
    "core_annotation": "核心注释字段",
    "cross_reference": "外部交叉引用",
    "evidence_source": "证据来源字段",
    "relationship_evidence": "关系证据字段",
    "activity_evidence": "活性/亲和力证据",
    "binding_site_evidence": "结合位点结构证据",
    "network_evidence": "网络/相互作用证据",
    "binding_site_annotation": "结合位点注释",
    "literature_reference": "文献引用字段",
    "clinical_status": "临床/批准状态",
    "derived_classification": "派生分类字段",
    "derived_interpretation": "派生解释字段",
    "derived_summary": "派生汇总字段",
    "annotation": "外部注释字段",
    "chemical_property": "化学性质字段",
    "chemical_structure": "化学结构字段",
    "low_coverage_safety": "低覆盖安全/毒性字段",
    "low_coverage_clinical": "低覆盖临床/药理字段",
    "pocket_geometry": "口袋几何字段",
    "reserved": "预留字段",
    "audit": "审计/溯源字段",
}


VALUE_MEANING_ZH = {
    "curated_target_relation": "来自 ChEMBL / DrugCentral 等整理过的靶点关系。",
    "bioassay_active_relation": "来自 PubChem BioAssay 的活性/阳性实验关系，不一定能定位结合位点。",
    "structure_ligand_context": "来自 BioLiP / sc-PDB / PDBbind 的同蛋白结构配体上下文。",
    "uniprot_binding_site_annotation": "来自 UniProt 实验结合位点注释中的明确 ligand。",
    "approved_drug": "当前数据捕获到已批准或已列名药物状态。",
    "clinical_trial_compound": "当前数据捕获到临床试验状态，但没有归为已批准药物。",
    "binding_evidence_only": "当前数据只捕获到结合、活性或结构证据，没有捕获到批准/临床状态。",
    "common_or_drug_name": "常用名或药物风格名称。",
    "research_code": "研发编号或实验化合物代码。",
    "systematic_or_chemical_name": "系统命名/IUPAC 风格化学名。",
    "pubchem_synonym_name": "从 PubChem 或映射来源选择的可读同义名。",
    "pdb_ligand_code_or_name": "主要来自 PDB ligand code 或结构配体上下文的名称。",
    "unknown_name_type": "无法明确判断名称类型。",
    "uniprot_ligand_name": "来自 UniProt experimental binding site 注释的 ligand 名称。",
    "therapeutic_or_clinical_compound": "治疗用或临床阶段化合物。",
    "bioactive_research_ligand": "有活性/结合证据的研究型小分子配体。",
    "structure_affinity_ligand": "具有结构和亲和力背景的配体，主要来自 PDBbind。",
    "endogenous_ligand_or_cofactor": "ATP、NAD、FMN、COA 等内源功能配体或辅因子。",
    "membrane_lipid_or_sterol": "膜脂、甾醇或脂质样内源配体。",
    "ion_or_metal": "简单离子或金属离子；当前规则较保守。",
    "buffer_salt_solvent": "缓冲液、盐、溶剂或结晶添加剂；当前规则较保守。",
    "biologic_or_large_molecule": "生物大分子或非典型小分子实体，保留原始 ID。",
    "unknown_structure_ligand": "结构中出现的 ligand，但目前无法可靠判断其生物角色。",
    "unmapped_uniprot_ligand": "来自 UniProt 实验结合位点注释，但当前没有可靠 PubChem CID 的 ligand。",
    "approved_or_clinical_target_ligand": "批准药物或临床化合物相关配体，用于治疗/靶点解释。",
    "curated_target_ligand": "来自整理过靶点关系的配体，例如 ChEMBL / DrugCentral。",
    "structure_bound_ligand": "结构中出现的配体，但没有被归为结构亲和力优先类。",
    "uniprot_experimental_site_ligand": "来自 UniProt 实验结合位点注释的明确非泛称 ligand。",
    "bioassay_active_ligand": "PubChem BioAssay 活性配体，有实验活性但结合位点可能不明确。",
    "cofactor_or_functional_ligand": "辅因子或功能性配体，可用于解释机制/功能位点。",
}


FIELD_ZH = {
    "target_uniprot_id": ("UniProt 蛋白 ID", "膜蛋白靶点的 UniProt accession，是蛋白相关表之间最重要的连接键。", "必需字段；理论上不应为空。", "来自源关系表和蛋白注释流程中的 UniProt 标准化。", "作为蛋白主键/外键使用，不建议用基因名替代。"),
    "approved_symbol": ("基因符号", "蛋白对应的 HGNC/常用基因符号，用于人类阅读。", "为空表示当前来源没有捕获到基因符号。", "来自 HGNC、UniProt 或源数据注释。", "用于展示和浏览，不作为主键。"),
    "ncbi_gene_id": ("NCBI Gene ID", "蛋白对应的 NCBI Gene 标识符；少数行可能用分号连接多个 ID。", "为空表示没有映射到 NCBI Gene ID。", "来自 NCBI/UniProt 映射缓存。", "用于和 NCBI/PubChem assay 体系交叉引用。"),
    "protein_name": ("蛋白名称", "蛋白的人类可读全名。", "为空表示当前来源没有捕获到蛋白名称。", "来自 UniProt 或源数据注释。", "用于展示和人工审阅。"),
    "drug_id": ("小分子/配体 ID", "化合物主标识符。通常是 PubChem CID；未映射 biologic 或结构 ligand 会保留稳定的非 CID ID。", "必需字段；理论上不应为空。", "CID 标准化流程；无法可靠转 CID 的对象保留原始或稳定占位 ID。", "作为小分子主键/外键使用。"),
    "compound_id_type": ("小分子 ID 类型", "说明 drug_id 属于哪个标识体系。", "必需分类字段；理论上不应为空。", "来自 CID 映射和结构 ligand 合并规则。", "需要 PubChem 工具时可筛选 `PubChem CID`。"),
    "drug_name": ("小分子名称", "小分子/配体的显示名称，可能来自源表、PubChem、PDB ligand 或原始名称。", "为空表示没有捕获到显示名，理论上应很少。", "源表、PubChem title/synonym、PDB ligand context 或原始名称。", "用于展示，不作为身份识别键。"),
    "compound_name_category": ("名称类型", "说明 drug_name 是常用名、研发代码、系统化学名、PubChem 同义名还是 PDB ligand 名。", "派生分类字段；理论上不应为空。", "由原 drug_name_type 和名称来源整理规则生成。", "用于区分可读药名、系统化学名和 PDB code 名称。"),
    "compound_biological_role": ("小分子生物角色", "描述这个化合物本身大体是什么角色，例如临床化合物、研究配体、结构亲和力配体、辅因子、脂质或未知结构 ligand。", "派生分类字段；理论上不应为空。", "由化合物状态、来源、PDBbind/结构上下文和保守角色规则生成。", "用于粗筛不同类型的小分子。"),
    "ligand_interpretation_class": ("配体解释类别", "binary 表行级别字段，说明这一条小分子-蛋白关系在靶点/结合位点解释上为什么有用。", "派生字段；理论上不应为空。", "由 compound_biological_role、compound_source、source_database 和临床/批准状态推导。", "作为行级别筛选证据价值和解释用途的核心字段。"),
    "ligand_interpretation_classes": ("配体解释类别合集", "小分子在 binary 表所有关系行中出现过的解释类别合集，用分号连接。", "派生汇总字段；理论上不应为空。", "按 drug_id 聚合 binary.ligand_interpretation_class。", "用于保留同一小分子的所有解释价值。"),
    "best_ligand_interpretation_class": ("最佳配体解释类别", "从 ligand_interpretation_classes 中按优先级选出的单一主类别。", "派生汇总字段；理论上不应为空。", "优先级：结构亲和力 > 批准/临床 > 辅因子 > 膜脂/甾醇 > curated target > structure-bound > bioassay active > ion/metal > buffer/salt/solvent > unknown。", "用于小分子层面的快速筛选。"),
    "compound_source": ("关系证据来源类型", "概括小分子-蛋白关系来自哪类证据来源。", "必需来源分类；理论上不应为空。", "由 source_database 和结构 ligand 合并规则推导。", "用于区分 curated target、bioassay active 和 structure ligand context。"),
    "source_database": ("来源数据库", "贡献当前 binary 关系行的数据库，可为多个来源合并。", "为空表示未捕获来源数据库，理论上应很少。", "来自原始数据表和合并流程。", "用于来源过滤和证据溯源。"),
    "source_databases": ("汇总来源数据库", "蛋白或小分子层面聚合后的来源数据库列表。", "为空表示该实体没有捕获到关系来源。", "从 binary.source_database 聚合得到。", "实体层面的来源概览；精确证据以 binary 表为准。"),
    "assay_or_mechanism": ("实验/机制描述", "实验 assay 描述、作用机制文字，或结构上下文描述。", "为空表示没有捕获到实验或机制文字。", "来自 ChEMBL、DrugCentral、PubChem BioAssay 或 PDB 上下文字段。", "用于人工审阅关系证据。"),
    "activity_type": ("活性类型", "实验活性类型，如 IC50、Ki、Kd、EC50 等。", "为空表示该行没有捕获到定量活性类型。", "来自源 assay/亲和力记录。", "需结合 activity_value_uM 使用，并注意不同 assay 不可简单横向比较。"),
    "activity_value_uM": ("活性值 uM", "归一化到微摩尔的活性/亲和力数值。", "为空表示没有捕获到数值活性。", "由源活性/亲和力数值单位换算得到。", "数值越低通常表示活性越强，但必须结合 assay 背景解释。"),
    "clinical_or_approval_status": ("行级临床/批准状态", "binary 行级别捕获到的批准或临床状态。", "为空表示该关系行没有捕获到临床/批准状态。", "来自药物/临床来源字段。", "小分子层面的筛选优先看 compound_evidence_status。"),
    "compound_evidence_status": ("小分子证据状态", "小分子层面区分已批准药物、临床化合物、仅有结合/活性/结构证据。", "派生字段；理论上不应为空。", "由 binary 表的 clinical_or_approval_status 聚合推导。", "用于区分正式药物/临床化合物和仅有结合证据的小分子。"),
    "has_pdb_biolip_matched": ("是否有 BioLiP 证据", "布尔标记，表示该行是否有 BioLiP 结构上下文/位点证据。", "理论上不应为空；0 表示当前数据未捕获。", "BioLiP 匹配/结构上下文合并。", "快速筛选 BioLiP 结构证据。"),
    "pdb_biolip_matched_sites": ("BioLiP 匹配位点", "BioLiP ligand 和残基/位点详情。", "为空表示该行没有 BioLiP 位点详情。", "BioLiP 解析后的位点字符串。", "用于残基级结构证据审阅。"),
    "has_pdb_scpdb_matched": ("是否有 sc-PDB 证据", "布尔标记，表示该行是否有 sc-PDB 结构上下文/位点证据。", "理论上不应为空；0 表示当前数据未捕获。", "sc-PDB 匹配/结构上下文合并。", "快速筛选 sc-PDB 结构证据。"),
    "pdb_scpdb_matched_sites": ("sc-PDB 匹配位点", "sc-PDB ligand 和残基/位点详情。", "为空表示该行没有 sc-PDB 位点详情。", "sc-PDB 解析后的位点字符串。", "用于残基级结构证据审阅。"),
    "has_pdbbind_matched": ("是否有 PDBbind 证据", "布尔标记，表示该行是否有 PDBbind 结构/亲和力上下文。", "理论上不应为空；0 表示当前数据未捕获。", "PDBbind 匹配/结构上下文合并。", "用于寻找结构-亲和力证据行。"),
    "pdbbind_matched_sites": ("PDBbind 匹配位点", "PDBbind ligand、位点和亲和力详情。", "为空表示该行没有 PDBbind 位点详情。", "PDBbind 复合物和亲和力字符串解析。", "对结合位点解释价值较高。"),
    "has_stitch_matched": ("是否有 STITCH 证据", "布尔标记，表示该行是否有 STITCH 高置信 compound-protein 证据。", "理论上不应为空；0 表示当前数据未捕获。", "清洗后的 STITCH 高置信 CID-UniProt 对。", "作为相互作用支持证据，不是直接结合位点证据。"),
    "stitch_matched_compounds": ("STITCH 匹配化合物", "CID-only 的 STITCH 证据字符串，格式为 CID:score。", "为空表示该行/蛋白没有保留的 STITCH 证据。", "清洗后的 STITCH 条目，已去掉 ChEMBL 前缀。", "作为支持性相互作用证据。"),
    "has_exp_binding_site": ("是否有 UniProt 实验位点", "布尔标记，表示该蛋白是否有 UniProt 文献/实验结合位点注释。", "理论上不应为空；0 表示当前数据未捕获。", "UniProt binding-site 注释。", "这是蛋白层面的位点背景，不一定是精确 drug-pair 证据。"),
    "exp_binding_sites": ("实验结合位点", "UniProt 残基级结合位点描述。", "为空表示未捕获 UniProt 实验结合位点。", "UniProt feature 注释。", "用于蛋白功能位点背景。"),
    "exp_ligands": ("实验注释配体", "UniProt 实验结合位点注释中提取的 ligand 名称。", "为空表示未捕获 UniProt experimental ligand。", "UniProt feature 注释。", "用于蛋白层面的已知配体背景。"),
    "exp_evidence_quality": ("实验位点证据质量", "UniProt binding-site 注释的证据质量等级。", "为空表示没有捕获证据质量。", "UniProt/ECO 证据映射。", "用于优先查看文献支持的结合位点。"),
    "exp_pubmed": ("实验位点 PubMed", "支持 UniProt 实验结合位点注释的 PubMed ID。", "为空表示没有捕获 PubMed ID。", "UniProt 注释。", "用于文献溯源。"),
    "exp_pdb": ("实验位点 PDB 引用", "预留给实验结合位点相关 PDB ID 的字段。", "当前版本为空；预留未来补充。", "预留的 UniProt/PDB 交叉引用字段。", "当前分析中不建议使用。"),
    "target_group_id": ("靶点组 ID", "保留原始多 UniProt 靶点记录拆分后的组信息。", "为空表示该行不是由多靶点源记录拆分而来。", "multi-UniProt explosion 过程中生成。", "只用于审计和追溯原始分组。"),
    "n_pockets_alphafold": ("AlphaFold 口袋数", "该蛋白的 AlphaFold 预测口袋实例数量。", "理论上不应为空；0 表示当前 pocket 表没有对应实例。", "从 pocket_instances.tsv 聚合得到。", "用于快速筛选有无预测口袋。"),
    "stitch_compounds_original_count": ("原始 STITCH 数量", "STITCH 高置信清洗前的原始条目数量。", "理论上不应为空；0 表示没有原始 STITCH 条目。", "STITCH 预处理。", "仅作为 STITCH 清洗背景。"),
    "stitch_compounds_cleaned": ("清洗后 STITCH 化合物", "保留的高置信 STITCH CID:score 条目。", "为空表示没有保留的 STITCH 条目。", "STITCH 高置信过滤。", "蛋白层面的相互作用支持信息。"),
    "stitch_compounds_cleaned_count": ("清洗后 STITCH 数量", "清洗后 STITCH 条目的数量。", "理论上不应为空；0 表示没有保留条目。", "stitch_compounds_cleaned 的计数。", "快速筛选 STITCH 覆盖。"),
    "unique_drug_count": ("唯一小分子数", "binary 表中与该蛋白相连的不同 drug_id 数量。", "理论上不应为空；0 表示没有关系行。", "从 binary 表重新计算。", "衡量蛋白关联小分子数量。"),
    "unique_drug_cid_sample": ("小分子 ID 示例", "该蛋白关联 drug_id 的最多 50 个预览样本。", "为空表示没有关系行。", "从 binary 表 drug_id 抽样。", "仅作预览；完整关系以 binary 表为准。"),
    "has_matched_evidence": ("是否有匹配证据", "蛋白层面是否至少有一种结构/STITCH/实验位点证据。", "理论上不应为空；0 表示当前没有捕获到匹配证据。", "由 binary 证据标记聚合得到。", "用于快速筛选有证据支持的蛋白。"),
    "reviewed": ("UniProt reviewed 状态", "UniProt/Swiss-Prot reviewed 状态。", "为空通常表示未审阅或未捕获。", "UniProt enrichment。", "用于优先查看 Swiss-Prot reviewed 蛋白。"),
    "go_cellular_component": ("GO 细胞组分", "GO cellular component 注释。", "为空表示没有捕获到 GO 细胞组分。", "UniProt/GO enrichment。", "用于定位/细胞组分背景。"),
    "go_molecular_function": ("GO 分子功能", "GO molecular function 注释。", "为空表示没有捕获到 GO 分子功能。", "UniProt/GO enrichment。", "用于功能背景。"),
    "pdb_structures": ("PDB 结构数量", "UniProt enrichment 中链接的 PDB 结构数量。", "理论上不应为空；0 表示未捕获到链接结构。", "UniProt xref_pdb enrichment。", "估计已有结构覆盖情况。"),
    "disease_association": ("疾病关联", "UniProt 疾病关联注释文本。", "为空表示没有捕获到疾病关联，不代表没有疾病相关性。", "UniProt disease comments。", "用于生物学背景，不作为完整疾病证据。"),
    "transmembrane_count": ("跨膜区数量", "UniProt TRANSMEM feature 数量。", "理论上不应为空；0 表示未计数到跨膜区。", "UniProt TRANSMEM feature 计数。", "用于表征膜蛋白拓扑。"),
    "subcellular_location": ("亚细胞定位", "UniProt 亚细胞定位注释。", "为空表示未捕获定位注释。", "UniProt subcellular location comments。", "用于定位背景。"),
    "protein_function": ("蛋白功能", "UniProt 功能注释文本。", "为空表示未捕获功能注释。", "UniProt function comments。", "用于蛋白功能解释。"),
    "source_row_count": ("关系行数", "该小分子在 binary 表中出现的关系/证据行数量。", "理论上不应为空；0 表示没有关系行。", "按 drug_id 从 binary 表聚合。", "这是关系行数，不是唯一靶点数。"),
    "unique_target_count": ("唯一靶点数", "该小分子连接的不同 target_uniprot_id 数量。", "理论上不应为空；0 表示没有连接蛋白。", "按 drug_id 从 binary 表聚合。", "用于衡量多靶点/广谱活性。"),
    "activity_types": ("活性类型汇总", "小分子层面观察到的活性类型集合。", "为空表示该小分子没有捕获到活性类型。", "由 binary.activity_type 聚合。", "用于了解该小分子有哪些定量证据。"),
    "activity_value_uM_min": ("最小活性值 uM", "该小分子所有关系中的最小活性/亲和力值。", "为空表示没有捕获到数值活性。", "由 binary.activity_value_uM 聚合。", "快速查看潜在强活性，但需注意 assay 背景。"),
    "activity_value_uM_median": ("中位活性值 uM", "该小分子所有关系中活性/亲和力值的中位数。", "为空表示没有捕获到数值活性。", "由 binary.activity_value_uM 聚合。", "比最小值更稳健，但仍需注意 assay 背景。"),
    "molecular_formula": ("分子式", "PubChem 分子式。", "为空通常表示没有 PubChem 属性，常见于未映射/biologic 条目。", "PubChem PUG REST 属性缓存。", "用于化学过滤和身份核对。"),
    "molecular_weight": ("分子量", "PubChem 分子量，单位 g/mol。", "为空表示没有取到 PubChem 属性。", "PubChem PUG REST 属性缓存。", "用于 drug-likeness 和 ligand 大小筛选。"),
    "canonical_smiles": ("Canonical SMILES", "PubChem canonical SMILES。", "为空表示没有取到 PubChem 结构。", "PubChem PUG REST 属性缓存。", "用于不强调立体化学的 cheminformatics。"),
    "isomeric_smiles": ("Isomeric SMILES", "包含立体化学信息的 PubChem isomeric SMILES。", "为空表示没有取到 PubChem 结构。", "PubChem PUG REST 属性缓存。", "立体化学重要时优先使用。"),
    "inchikey": ("InChIKey", "PubChem InChIKey。", "为空表示没有可用化学结构映射。", "PubChem PUG REST 属性缓存。", "用于跨数据库化学身份匹配。"),
    "iupac_name": ("IUPAC 名称", "PubChem IUPAC 系统命名。", "为空表示没有取到 IUPAC 名称。", "PubChem PUG REST 属性缓存。", "用于精确化学命名，可能很长。"),
    "xlogp": ("XLogP", "PubChem XLogP 脂溶性估计值。", "为空表示没有计算/取到 XLogP。", "PubChem PUG REST 属性缓存。", "用于疏水性和膜通透性启发式筛选。"),
    "tpsa": ("TPSA", "拓扑极性表面积。", "为空表示没有取到属性。", "PubChem PUG REST 属性缓存。", "用于通透性/drug-likeness 筛选。"),
    "hbond_donor_count": ("氢键供体数", "PubChem 氢键供体数量。", "为空表示没有取到属性。", "PubChem PUG REST 属性缓存。", "用于 drug-likeness 筛选。"),
    "hbond_acceptor_count": ("氢键受体数", "PubChem 氢键受体数量。", "为空表示没有取到属性。", "PubChem PUG REST 属性缓存。", "用于 drug-likeness 筛选。"),
    "rotatable_bond_count": ("可旋转键数", "PubChem 可旋转键数量。", "为空表示没有取到属性。", "PubChem PUG REST 属性缓存。", "用于柔性和 drug-likeness 筛选。"),
    "complexity": ("分子复杂度", "PubChem molecular complexity 分数。", "为空表示没有取到属性。", "PubChem PUG REST 属性缓存。", "用于粗略结构复杂度评估。"),
    "formal_charge": ("形式电荷", "PubChem formal charge。", "为空表示没有取到属性。", "PubChem PUG REST 属性缓存。", "用于识别带电化合物/离子。"),
    "ghs_hazard_classes": ("GHS 危害类别", "PubChem/PUG-View 捕获到的 GHS hazard class 文本。", "`not_found_in_current_sources` 表示当前来源没有提供该字段，不代表没有危害。", "PubChem PUG-View 安全/毒性提取。", "作为稀疏安全注释谨慎使用。"),
    "ghs_signal_words": ("GHS 警示词", "GHS signal word，例如 Warning 或 Danger。", "`not_found_in_current_sources` 表示当前来源没有提供。", "PubChem PUG-View 安全/毒性提取。", "当前覆盖很低，谨慎使用。"),
    "toxicity_summary": ("毒性摘要", "捕获到的人类可读毒性摘要。", "`not_found_in_current_sources` 表示当前来源没有提供毒性摘要，不代表没有毒性。", "PubChem PUG-View 安全/毒性提取。", "仅适合已知化合物的定性审阅。"),
    "livertox": ("LiverTox 摘要", "捕获到的 LiverTox 文本摘要。", "`not_found_in_current_sources` 表示没有捕获到 LiverTox 摘要。", "PubChem PUG-View / LiverTox 提取。", "用于已知药物肝毒性背景。"),
    "drug_classes": ("药物类别", "捕获到的药物类别/分类文本。", "`not_found_in_current_sources` 表示没有捕获到药物类别文本。", "PubChem PUG-View 临床/药理提取。", "作为稀疏临床注释使用。"),
    "pharmacodynamics": ("药效学描述", "捕获到的 pharmacodynamics 文本。", "`not_found_in_current_sources` 表示没有捕获到药效学文本。", "PubChem PUG-View 临床/药理提取。", "用于定性药理审阅。"),
    "drug_indication": ("药物适应症", "捕获到的适应症/用途描述。", "`not_found_in_current_sources` 表示没有捕获到适应症文本，不代表没有医学用途。", "PubChem PUG-View 临床/药理提取。", "作为稀疏临床注释使用。"),
    "pocket_batch": ("口袋预测批次", "AlphaFold pocket prediction 输入批次标签。", "理论上不应为空。", "AlphaFold 口袋预测批次元数据。", "只用于溯源/QC。"),
    "pocket_id": ("口袋 ID", "唯一 pocket identifier，通常由 UniProt 加 pocket 编号组成。", "必需 pocket key；理论上不应为空。", "AlphaFold 口袋预测流程。", "作为 pocket_instances 表主键使用。"),
    "center_x": ("口袋中心 X", "预测口袋中心的 X 坐标。", "理论上不应为空。", "AlphaFold 口袋预测流程。", "用于空间口袋分析。"),
    "center_y": ("口袋中心 Y", "预测口袋中心的 Y 坐标。", "理论上不应为空。", "AlphaFold 口袋预测流程。", "用于空间口袋分析。"),
    "center_z": ("口袋中心 Z", "预测口袋中心的 Z 坐标。", "理论上不应为空。", "AlphaFold 口袋预测流程。", "用于空间口袋分析。"),
    "n_residues": ("口袋残基数", "该预测口袋包含的残基数量。", "理论上不应为空。", "AlphaFold 口袋预测流程。", "用于筛选口袋大小。"),
    "score": ("口袋评分", "AlphaFold 口袋预测流程给出的启发式口袋评分。", "理论上不应为空。", "AlphaFold 口袋预测流程。", "可用于同蛋白内口袋排序，需谨慎解释。"),
    "pocket_residues": ("口袋残基列表", "预测口袋包含的残基列表，通常以分号分隔。", "理论上不应为空。", "AlphaFold 口袋预测流程。", "用于残基级口袋比较。"),
    "source_file": ("来源文件", "原始 pocket prediction CSV 文件名。", "理论上不应为空。", "口袋预测流程输出文件名。", "用于溯源/debug。"),
}


ROLE_MEANINGS_ZH = {
    "核心标识字段": "主键、外键或核心连接字段。",
    "核心注释字段": "名称、标签等基础可读信息。",
    "外部交叉引用": "指向外部数据库的 ID。",
    "证据来源字段": "说明关系或注释来自哪个数据源。",
    "关系证据字段": "小分子-蛋白关系层面的 assay 或机制证据。",
    "活性/亲和力证据": "定量活性、亲和力或 assay 数值字段。",
    "结合位点结构证据": "来自 PDB/BioLiP/sc-PDB/PDBbind 等结构/位点证据。",
    "网络/相互作用证据": "STITCH 等网络或相互作用证据。",
    "结合位点注释": "UniProt 等来源的蛋白层面结合位点注释。",
    "文献引用字段": "PubMed 等文献溯源字段。",
    "临床/批准状态": "临床试验或批准/列名相关字段。",
    "派生分类字段": "由规则从原始字段推导出的分类。",
    "派生解释字段": "用于解释配体/关系价值的派生字段。",
    "派生汇总字段": "从行级表聚合到实体表的汇总字段。",
    "外部注释字段": "UniProt/GO/PubChem 等外部注释。",
    "化学性质字段": "分子量、LogP、TPSA 等化学性质。",
    "化学结构字段": "SMILES、InChIKey 等结构表示。",
    "低覆盖安全/毒性字段": "覆盖率较低的安全和毒性文本字段。",
    "低覆盖临床/药理字段": "覆盖率较低的临床和药理文本字段。",
    "口袋几何字段": "AlphaFold 预测口袋的几何信息。",
    "预留字段": "当前未填充或为未来补充保留的字段。",
    "审计/溯源字段": "用于追踪来源、批次或处理过程的字段。",
}


def zh_bool(value: str) -> str:
    return {"yes": "是", "no": "否"}.get(value, value)


def translate_allowed_values(value: str) -> str:
    if not value:
        return ""
    parts = [part.strip() for part in value.split(";") if part.strip()]
    return "; ".join(f"{part} = {VALUE_MEANING_ZH.get(part, part)}" for part in parts)


def read_rows() -> list[dict[str, str]]:
    with SRC.open("r", encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f, delimiter="\t"))


def convert_rows(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    converted = []
    for row in rows:
        field = row["column_name"]
        label, desc, missing, source, recommended = FIELD_ZH.get(
            field,
            (
                field.replace("_", " "),
                f"`{field}` 字段，来自当前标准化数据表。",
                "为空表示当前来源没有捕获到该值。",
                "由标准化表生成流程产生。",
                "使用前建议结合字段角色和覆盖率判断。",
            ),
        )
        role = ROLE_ZH.get(row["field_role"], row["field_role"])
        notes = row["notes"]
        if "0% coverage" in notes:
            notes = "当前版本覆盖率为 0%。"
        elif "Not equivalent to drug approval status" in notes:
            notes = "该字段不是药物批准状态，而是本数据集中用于解释配体价值的分类。"
        elif "Consider renaming to relationship_row_count" in notes:
            notes = "未来可考虑改名为 relationship_row_count，避免误解为来源数量。"
        converted.append(
            {
                "所属Sheet": row["sheet_name"],
                "对应TSV文件": row["table_file"],
                "列序号": row["column_order"],
                "字段名": field,
                "中文字段名": label,
                "字段说明": desc,
                "数据类型": TYPE_ZH.get(row["data_type"], row["data_type"]),
                "是否主键": zh_bool(row["is_primary_key"]),
                "是否外键": zh_bool(row["is_foreign_key"]),
                "外键指向": row["references"],
                "允许值": translate_allowed_values(row["allowed_values"]),
                "多值分隔符": row["separator"],
                "缺失值含义": missing,
                "字段角色": role,
                "来源或生成规则": source,
                "总行数": row["total_rows"],
                "真实有值数": row["real_value_count"],
                "not_found计数": row["not_found_count"],
                "空值数": row["empty_count"],
                "真实覆盖率%": row["coverage_percent"],
                "观察到的唯一值数": row["observed_unique_count"],
                "高频取值": row["top_values"],
                "示例值": row["example_values"],
                "备注": notes,
                "推荐用法": recommended,
            }
        )
    return converted


def allowed_value_rows() -> list[dict[str, str]]:
    rows = []
    for raw in read_rows():
        field = raw["column_name"]
        allowed = [part.strip() for part in raw["allowed_values"].split(";") if part.strip()]
        for value in allowed:
            rows.append(
                {
                    "字段名": field,
                    "允许值": value,
                    "中文含义": VALUE_MEANING_ZH.get(value, ""),
                }
            )
    seen = set()
    unique = []
    for row in rows:
        key = (row["字段名"], row["允许值"])
        if key not in seen:
            unique.append(row)
            seen.add(key)
    return unique


def write_tsv(rows: list[dict[str, str]]) -> None:
    with OUT_TSV.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=HEADERS_ZH, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(rows: list[dict[str, str]]) -> None:
    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")

    ws = wb.active
    ws.title = "字段字典"
    ws.append(HEADERS_ZH)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        ws.append([row[h] for h in HEADERS_ZH])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS_ZH))}{len(rows)+1}"
    widths = [24, 34, 10, 32, 28, 56, 14, 10, 10, 38, 70, 12, 56, 24, 58, 12, 14, 14, 10, 12, 16, 64, 64, 52, 56]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws2 = wb.create_sheet("允许值说明")
    allowed_headers = ["字段名", "允许值", "中文含义"]
    ws2.append(allowed_headers)
    for cell in ws2[1]:
        cell.fill = header_fill
        cell.font = header_font
    for row in allowed_value_rows():
        ws2.append([row[h] for h in allowed_headers])
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:C{ws2.max_row}"
    ws2.column_dimensions["A"].width = 38
    ws2.column_dimensions["B"].width = 42
    ws2.column_dimensions["C"].width = 90
    for row in ws2.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws3 = wb.create_sheet("字段角色说明")
    role_headers = ["字段角色", "字段数量", "中文含义"]
    ws3.append(role_headers)
    counts: dict[str, int] = {}
    for row in rows:
        counts[row["字段角色"]] = counts.get(row["字段角色"], 0) + 1
    for role, count in sorted(counts.items()):
        ws3.append([role, count, ROLE_MEANINGS_ZH.get(role, "")])
    for cell in ws3[1]:
        cell.fill = header_fill
        cell.font = header_font
    ws3.freeze_panes = "A2"
    ws3.auto_filter.ref = f"A1:C{ws3.max_row}"
    ws3.column_dimensions["A"].width = 30
    ws3.column_dimensions["B"].width = 12
    ws3.column_dimensions["C"].width = 90
    for row in ws3.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(OUT_XLSX)


def main() -> None:
    rows = convert_rows(read_rows())
    write_tsv(rows)
    write_xlsx(rows)
    print({"rows": len(rows), "tsv": str(OUT_TSV), "xlsx": str(OUT_XLSX)})


if __name__ == "__main__":
    main()
