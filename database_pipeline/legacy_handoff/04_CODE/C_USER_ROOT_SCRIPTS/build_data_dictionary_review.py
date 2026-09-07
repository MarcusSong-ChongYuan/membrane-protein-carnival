from __future__ import annotations

import csv
import json
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


BASE = Path(r"C:\github-repos\upload\normalized_tables")
OUT_TSV = BASE / "data_dictionary_review.tsv"
OUT_XLSX = BASE / "data_dictionary_review.xlsx"

NOT_FOUND = "not_found_in_current_sources"

TABLES = [
    ("1_Binary_Relationships", "drug_protein_binary_relationships.tsv"),
    ("2_Protein_Database", "protein_database.tsv"),
    ("3_Small_Molecules", "small_molecule_database.tsv"),
    ("4_Pocket_Instances", "pocket_instances.tsv"),
]

DICT_COLUMNS = [
    "sheet_name",
    "table_file",
    "column_order",
    "column_name",
    "column_label",
    "description",
    "data_type",
    "is_primary_key",
    "is_foreign_key",
    "references",
    "allowed_values",
    "separator",
    "missing_value_meaning",
    "field_role",
    "source_or_rule",
    "total_rows",
    "real_value_count",
    "not_found_count",
    "empty_count",
    "coverage_percent",
    "observed_unique_count",
    "top_values",
    "example_values",
    "notes",
    "recommended_use",
]

ALLOWED_VALUE_NOTES = {
    "compound_source": {
        "curated_target_relation": "Curated target relationship from ChEMBL/DrugCentral-style sources.",
        "bioassay_active_relation": "Active/positive PubChem BioAssay relationship; may not localize the binding site.",
        "structure_ligand_context": "Ligand observed in BioLiP/sc-PDB/PDBbind same-protein structural context.",
        "uniprot_binding_site_annotation": "Explicit ligand from UniProt experimental binding-site annotation.",
    },
    "compound_evidence_status": {
        "approved_drug": "Approved/listed status captured for this compound.",
        "clinical_trial_compound": "Clinical-trial status captured, but approval/listed status was not captured.",
        "binding_evidence_only": "Current dataset only captured binding/activity/structure evidence.",
    },
    "compound_name_category": {
        "common_or_drug_name": "Readable common/drug-style name.",
        "research_code": "Development or research code.",
        "systematic_or_chemical_name": "Systematic/IUPAC-like chemical name.",
        "pubchem_synonym_name": "Readable synonym selected from PubChem or mapping sources.",
        "pdb_ligand_code_or_name": "Name mainly derived from PDB ligand code/context.",
        "unknown_name_type": "Name category could not be determined.",
    },
    "compound_biological_role": {
        "therapeutic_or_clinical_compound": "Approved or clinical compound.",
        "bioactive_research_ligand": "Research/bioassay ligand with activity/binding evidence.",
        "structure_affinity_ligand": "Ligand with structure plus affinity context, mainly PDBbind.",
        "endogenous_ligand_or_cofactor": "ATP/NAD/FMN/COA-like functional ligand or cofactor.",
        "membrane_lipid_or_sterol": "Membrane lipid, sterol, or lipid-like endogenous ligand.",
        "ion_or_metal": "Simple ion or metal; currently conservatively assigned.",
        "buffer_salt_solvent": "Buffer, salt, solvent, or crystallization additive; currently conservatively assigned.",
        "biologic_or_large_molecule": "Biologic or large non-small-molecule entity retained with original ID.",
        "unknown_structure_ligand": "Structure ligand whose biological role is not yet clear.",
    },
    "ligand_interpretation_class": {
        "structure_affinity_ligand": "Row has direct structure/affinity interpretation value, usually PDBbind.",
        "approved_or_clinical_target_ligand": "Approved/clinical ligand relationship; useful for therapeutic interpretation.",
        "cofactor_or_functional_ligand": "Functional ligand/cofactor; useful for mechanistic site interpretation.",
        "membrane_lipid_or_sterol": "Membrane lipid/sterol; particularly relevant to membrane proteins.",
        "uniprot_experimental_site_ligand": "Explicit non-generic ligand from UniProt experimental binding-site annotation.",
        "curated_target_ligand": "Curated target relation from ChEMBL/DrugCentral-style sources.",
        "structure_bound_ligand": "Structure-context ligand without affinity-priority classification.",
        "bioassay_active_ligand": "BioAssay-active ligand; activity evidence but binding site may be unclear.",
        "ion_or_metal": "Ion/metal relationship; useful but should be interpreted separately.",
        "buffer_salt_solvent": "Likely buffer/salt/solvent/additive; usually not a priority ligand.",
        "unknown_structure_ligand": "Structure ligand retained but not yet interpretable.",
    },
}
ALLOWED_VALUE_NOTES["ligand_interpretation_classes"] = ALLOWED_VALUE_NOTES["ligand_interpretation_class"]
ALLOWED_VALUE_NOTES["best_ligand_interpretation_class"] = ALLOWED_VALUE_NOTES["ligand_interpretation_class"]

BASE_META: dict[str, dict[str, str]] = {
    "target_uniprot_id": {
        "label": "UniProt protein ID",
        "description": "Primary UniProt accession for the membrane protein target.",
        "type": "identifier",
        "missing": "Required key; should not be empty.",
        "role": "core_identifier",
        "source": "Mapped/standardized UniProt accession from source relationship and protein annotation pipelines.",
        "recommended": "Use as the main protein join key across binary, protein, and pocket tables.",
    },
    "approved_symbol": {
        "label": "Gene symbol",
        "description": "Approved/readable gene symbol for the protein.",
        "type": "text",
        "missing": "Empty means no approved symbol was captured.",
        "role": "core_annotation",
        "source": "HGNC/UniProt/source annotation.",
        "recommended": "Use for display; do not use as primary key.",
    },
    "ncbi_gene_id": {
        "label": "NCBI Gene ID",
        "description": "NCBI Gene identifier for the protein. Some rows may contain multiple IDs separated by semicolon.",
        "type": "identifier",
        "separator": ";",
        "missing": "Empty means no NCBI Gene ID was mapped.",
        "role": "cross_reference",
        "source": "NCBI/UniProt mapping cache.",
        "recommended": "Use for NCBI/PubChem assay cross-reference; prefer UniProt for joins.",
    },
    "protein_name": {
        "label": "Protein name",
        "description": "Human-readable protein name.",
        "type": "text",
        "missing": "Empty means no protein name was captured.",
        "role": "core_annotation",
        "source": "UniProt/source annotation.",
        "recommended": "Use for browsing and display.",
    },
    "drug_id": {
        "label": "Compound ID",
        "description": "Primary compound identifier. Usually PubChem CID; biologics and unmapped structure ligands may retain stable non-CID IDs.",
        "type": "identifier",
        "missing": "Required key; should not be empty.",
        "role": "core_identifier",
        "source": "CID normalization pipeline; unmapped IDs retained when CID conversion was not reliable.",
        "recommended": "Use as the main compound join key between binary and small-molecule tables.",
    },
    "compound_id_type": {
        "label": "Compound ID type",
        "description": "Identifier namespace/type for drug_id.",
        "type": "category",
        "allowed": "PubChem CID;ChEMBL;DrugCentral;unmapped_structure_ligand;unmapped_uniprot_ligand",
        "missing": "Required classification; should not be empty.",
        "role": "core_identifier",
        "source": "CID mapping and structure-ligand merge rules.",
        "recommended": "Filter to PubChem CID when downstream tools require PubChem identifiers.",
    },
    "drug_name": {
        "label": "Compound name",
        "description": "Preferred display name for the compound/ligand.",
        "type": "text",
        "missing": "Empty means no display name was captured; should be rare.",
        "role": "core_annotation",
        "source": "Source tables, PubChem title/synonym lookup, PDB ligand context, or retained original name.",
        "recommended": "Use for display, not as identity key.",
    },
    "compound_name_category": {
        "label": "Compound name category",
        "description": "Classification of what kind of display name is stored in drug_name.",
        "type": "category",
        "allowed": ";".join(list(ALLOWED_VALUE_NOTES["compound_name_category"].keys()) + ["uniprot_ligand_name"]),
        "missing": "Required derived category; should not be empty.",
        "role": "derived_classification",
        "source": "Derived from original drug_name_type/name-source cleanup rules.",
        "recommended": "Use to distinguish drug/common names from systematic names or PDB ligand-code names.",
    },
    "compound_biological_role": {
        "label": "Compound biological role",
        "description": "Broad biological or practical role of the compound itself.",
        "type": "category",
        "allowed": ";".join(ALLOWED_VALUE_NOTES["compound_biological_role"].keys()),
        "missing": "Required derived category; should not be empty.",
        "role": "derived_classification",
        "source": "Derived from compound status, source, PDBbind/structure context, and conservative role rules.",
        "recommended": "Use for broad filtering of therapeutic compounds, bioactive research ligands, cofactors, lipids, biologics, and unknown structure ligands.",
    },
    "ligand_interpretation_class": {
        "label": "Row-level ligand interpretation class",
        "description": "Why this ligand-protein row is useful for target/binding-site interpretation.",
        "type": "category",
        "allowed": ";".join(ALLOWED_VALUE_NOTES["ligand_interpretation_class"].keys()),
        "missing": "Required derived row-level class; should not be empty.",
        "role": "derived_interpretation",
        "source": "Derived from compound_biological_role, compound_source, source_database, and clinical/approval status.",
        "recommended": "Use this as the main row-level filter for ligand evidence quality and interpretation purpose.",
        "notes": "Not equivalent to drug approval status; it is an interpretation class for this dataset.",
    },
    "ligand_interpretation_classes": {
        "label": "All ligand interpretation classes",
        "description": "All row-level ligand interpretation classes observed for this compound across the binary table.",
        "type": "multi_category",
        "allowed": ";".join(ALLOWED_VALUE_NOTES["ligand_interpretation_class"].keys()),
        "separator": ";",
        "missing": "Required derived compound-level summary; should not be empty.",
        "role": "derived_summary",
        "source": "Aggregated from binary.ligand_interpretation_class by drug_id.",
        "recommended": "Use when you need all reasons a compound may matter.",
    },
    "best_ligand_interpretation_class": {
        "label": "Best ligand interpretation class",
        "description": "Single highest-priority ligand interpretation class for a compound.",
        "type": "category",
        "allowed": ";".join(ALLOWED_VALUE_NOTES["ligand_interpretation_class"].keys()),
        "missing": "Required derived compound-level class; should not be empty.",
        "role": "derived_summary",
        "source": "Selected from ligand_interpretation_classes using priority: structure affinity > approved/clinical > cofactor > membrane lipid/sterol > curated target > structure-bound > bioassay active > ion/metal > buffer/salt/solvent > unknown structure ligand.",
        "recommended": "Use for fast compound-level filtering.",
    },
    "compound_source": {
        "label": "Relationship evidence origin",
        "description": "Broad origin of the compound-protein relationship evidence.",
        "type": "multi_category",
        "allowed": ";".join(ALLOWED_VALUE_NOTES["compound_source"].keys()),
        "separator": ";",
        "missing": "Required source class; should not be empty.",
        "role": "evidence_source",
        "source": "Derived from source_database and structure-ligand merge rules.",
        "recommended": "Use to separate curated target relations, active bioassays, and structure ligand context.",
    },
    "source_database": {
        "label": "Source database",
        "description": "Database(s) contributing this binary relationship row.",
        "type": "multi_category",
        "separator": ";",
        "missing": "Empty means source database was not captured; should be rare.",
        "role": "evidence_source",
        "source": "Original source tables and merge pipeline.",
        "recommended": "Use for source-specific filtering or provenance.",
    },
    "source_databases": {
        "label": "Aggregated source databases",
        "description": "Semicolon-separated databases contributing relationships for this protein or compound.",
        "type": "multi_category",
        "separator": ";",
        "missing": "Empty means no relationship source was captured for the entity.",
        "role": "derived_summary",
        "source": "Aggregated from binary.source_database.",
        "recommended": "Use for entity-level source provenance; use binary rows for exact evidence.",
    },
    "assay_or_mechanism": {
        "label": "Assay or mechanism",
        "description": "Assay description, mechanism-of-action text, or structure-context descriptor.",
        "type": "text",
        "missing": "Empty means no assay/mechanism text was captured.",
        "role": "relationship_evidence",
        "source": "ChEMBL/DrugCentral/PubChem BioAssay/PDB context source fields.",
        "recommended": "Use for manual review of relationship evidence.",
    },
    "activity_type": {
        "label": "Activity type",
        "description": "Experimental activity type such as IC50, Ki, Kd, EC50, or related labels.",
        "type": "category",
        "separator": ";",
        "missing": "Empty means no quantitative activity type was captured for this row.",
        "role": "activity_evidence",
        "source": "Source assay/affinity records.",
        "recommended": "Use together with activity_value_uM; compare only compatible assay types.",
    },
    "activity_value_uM": {
        "label": "Activity value (uM)",
        "description": "Activity/affinity value normalized to micromolar when available.",
        "type": "float",
        "missing": "Empty means no numeric activity value was captured.",
        "role": "activity_evidence",
        "source": "Normalized from source activity/affinity values.",
        "recommended": "Lower values generally indicate stronger measured activity, but assay context matters.",
    },
    "clinical_or_approval_status": {
        "label": "Row clinical/approval status",
        "description": "Row-level clinical or approval status captured from relationship sources.",
        "type": "multi_category",
        "allowed": "approved_or_listed;clinical_trial",
        "separator": ";",
        "missing": "Empty means no row-level clinical/approval status was captured.",
        "role": "clinical_status",
        "source": "Drug/clinical source fields in the relationship pipeline.",
        "recommended": "Use compound_evidence_status in the small-molecule table for compound-level filtering.",
    },
    "compound_evidence_status": {
        "label": "Compound evidence status",
        "description": "Compound-level status distinguishing approved/listed, clinical-trial, and binding-evidence-only compounds.",
        "type": "category",
        "allowed": "approved_drug;clinical_trial_compound;binding_evidence_only",
        "missing": "Required derived status; should not be empty.",
        "role": "derived_classification",
        "source": "Derived from relationship clinical_or_approval_status.",
        "recommended": "Use to separate formal drug/clinical compounds from compounds with only binding/activity/structure evidence.",
    },
    "has_pdb_biolip_matched": {
        "label": "Has BioLiP evidence",
        "description": "Boolean flag indicating BioLiP structural context/evidence for this row.",
        "type": "boolean",
        "allowed": "0;1",
        "missing": "Should not be empty; 0 means absent in current data.",
        "role": "binding_site_evidence",
        "source": "BioLiP matching/structure-context merge.",
        "recommended": "Use for quick structural evidence filtering.",
    },
    "pdb_biolip_matched_sites": {
        "label": "BioLiP matched sites",
        "description": "BioLiP ligand and residue/site detail for this row.",
        "type": "text",
        "separator": ";",
        "missing": "Empty means no BioLiP site detail was captured for this row.",
        "role": "binding_site_evidence",
        "source": "BioLiP parsed site strings.",
        "recommended": "Use for residue-level structure evidence review.",
    },
    "has_pdb_scpdb_matched": {
        "label": "Has sc-PDB evidence",
        "description": "Boolean flag indicating sc-PDB structural context/evidence for this row.",
        "type": "boolean",
        "allowed": "0;1",
        "missing": "Should not be empty; 0 means absent in current data.",
        "role": "binding_site_evidence",
        "source": "sc-PDB matching/structure-context merge.",
        "recommended": "Use for quick structural evidence filtering.",
    },
    "pdb_scpdb_matched_sites": {
        "label": "sc-PDB matched sites",
        "description": "sc-PDB ligand and residue/site detail for this row.",
        "type": "text",
        "separator": ";",
        "missing": "Empty means no sc-PDB site detail was captured for this row.",
        "role": "binding_site_evidence",
        "source": "sc-PDB parsed site strings.",
        "recommended": "Use for residue-level structure evidence review.",
    },
    "has_pdbbind_matched": {
        "label": "Has PDBbind evidence",
        "description": "Boolean flag indicating PDBbind structure/affinity context for this row.",
        "type": "boolean",
        "allowed": "0;1",
        "missing": "Should not be empty; 0 means absent in current data.",
        "role": "binding_site_evidence",
        "source": "PDBbind matching/structure-context merge.",
        "recommended": "Use to find structure-affinity ligand rows.",
    },
    "pdbbind_matched_sites": {
        "label": "PDBbind matched sites",
        "description": "PDBbind ligand/site/affinity detail for this row.",
        "type": "text",
        "separator": ";",
        "missing": "Empty means no PDBbind site detail was captured for this row.",
        "role": "binding_site_evidence",
        "source": "PDBbind parsed complex and affinity strings.",
        "recommended": "High-value field for binding-site interpretation.",
    },
    "has_stitch_matched": {
        "label": "Has STITCH evidence",
        "description": "Boolean flag indicating STITCH high-confidence compound-protein evidence for this row.",
        "type": "boolean",
        "allowed": "0;1",
        "missing": "Should not be empty; 0 means absent in current data.",
        "role": "network_evidence",
        "source": "STITCH cleaned high-confidence CID-UniProt pairs.",
        "recommended": "Use as supporting evidence; not direct binding-site evidence.",
    },
    "stitch_matched_compounds": {
        "label": "STITCH matched compounds",
        "description": "CID-only STITCH evidence strings in the form CID:score.",
        "type": "multi_category",
        "separator": ";",
        "missing": "Empty means no retained STITCH evidence for this row/protein.",
        "role": "network_evidence",
        "source": "Cleaned STITCH entries; ChEMBL prefixes removed.",
        "recommended": "Use as supporting interaction evidence.",
    },
    "has_exp_binding_site": {
        "label": "Has UniProt experimental site",
        "description": "Boolean flag indicating UniProt literature/experimental binding site annotation exists for the protein.",
        "type": "boolean",
        "allowed": "0;1",
        "missing": "Should not be empty; 0 means no captured UniProt binding-site annotation.",
        "role": "binding_site_annotation",
        "source": "UniProt binding-site annotations.",
        "recommended": "Use as protein-level site annotation, not necessarily exact drug-pair evidence.",
    },
    "exp_binding_sites": {
        "label": "Experimental binding sites",
        "description": "UniProt residue-level binding-site descriptions.",
        "type": "text",
        "separator": ";",
        "missing": "Empty means no UniProt experimental binding site was captured.",
        "role": "binding_site_annotation",
        "source": "UniProt feature annotations.",
        "recommended": "Use for protein-level functional site context.",
    },
    "exp_ligands": {
        "label": "Experimental ligands",
        "description": "Ligand names extracted from UniProt experimental binding-site annotations.",
        "type": "multi_category",
        "separator": ";",
        "missing": "Empty means no UniProt experimental ligand was captured.",
        "role": "binding_site_annotation",
        "source": "UniProt feature annotations.",
        "recommended": "Use as protein-level known-ligand context.",
    },
    "exp_evidence_quality": {
        "label": "Experimental evidence quality",
        "description": "Evidence quality tiers for UniProt binding-site annotations.",
        "type": "multi_category",
        "separator": ";",
        "missing": "Empty means no UniProt binding-site evidence quality was captured.",
        "role": "binding_site_annotation",
        "source": "UniProt/ECO evidence mapping.",
        "recommended": "Use to prioritize literature-supported binding-site annotations.",
    },
    "exp_pubmed": {
        "label": "Experimental PubMed IDs",
        "description": "PubMed IDs supporting UniProt experimental binding-site annotations.",
        "type": "multi_identifier",
        "separator": ";",
        "missing": "Empty means no PubMed ID was captured for the annotation.",
        "role": "literature_reference",
        "source": "UniProt annotations.",
        "recommended": "Use for literature traceability.",
    },
    "exp_pdb": {
        "label": "Experimental PDB cross-reference",
        "description": "Reserved field for PDB IDs linked from experimental binding-site annotation.",
        "type": "multi_identifier",
        "separator": ";",
        "missing": "Currently empty in the dataset; reserved for future enrichment.",
        "role": "reserved",
        "source": "Reserved UniProt/PDB cross-reference field.",
        "recommended": "Do not use for current analysis until populated.",
        "notes": "0% coverage in current release.",
    },
    "target_group_id": {
        "label": "Target group ID",
        "description": "Group ID preserving source rows that originally listed multiple UniProt targets.",
        "type": "identifier",
        "missing": "Empty means the row was not generated from a multi-target source record.",
        "role": "audit",
        "source": "Generated during multi-UniProt explosion.",
        "recommended": "Use only for traceability to original grouped target records.",
    },
    "n_pockets_alphafold": {
        "label": "AlphaFold pocket count",
        "description": "Number of predicted AlphaFold pocket instances for the protein.",
        "type": "integer",
        "missing": "Should not be empty; 0 means no pocket instance in current pocket table.",
        "role": "derived_summary",
        "source": "Aggregated from pocket_instances.tsv.",
        "recommended": "Use for quick protein-level pocket availability filtering.",
    },
    "stitch_compounds_original_count": {
        "label": "Original STITCH count",
        "description": "Number of original STITCH entries before high-confidence filtering.",
        "type": "integer",
        "missing": "Should not be empty; 0 means no original STITCH compounds captured.",
        "role": "derived_summary",
        "source": "STITCH preprocessing.",
        "recommended": "Use only as context for STITCH filtering.",
    },
    "stitch_compounds_cleaned": {
        "label": "Cleaned STITCH compounds",
        "description": "Retained high-confidence STITCH CID:score entries.",
        "type": "multi_category",
        "separator": ";",
        "missing": "Empty means no retained cleaned STITCH entries.",
        "role": "network_evidence",
        "source": "STITCH cleaned high-confidence filter.",
        "recommended": "Use as protein-level supporting interaction context.",
    },
    "stitch_compounds_cleaned_count": {
        "label": "Cleaned STITCH count",
        "description": "Number of retained cleaned STITCH entries.",
        "type": "integer",
        "missing": "Should not be empty; 0 means no retained cleaned STITCH entries.",
        "role": "derived_summary",
        "source": "Count of stitch_compounds_cleaned entries.",
        "recommended": "Use for quick STITCH coverage filtering.",
    },
    "unique_drug_count": {
        "label": "Unique compound count",
        "description": "Number of distinct drug_id values linked to this protein in the binary table.",
        "type": "integer",
        "missing": "Should not be empty; 0 means no binary compound relationship for this protein.",
        "role": "derived_summary",
        "source": "Recomputed from binary table.",
        "recommended": "Use for protein-level relationship density.",
    },
    "unique_drug_cid_sample": {
        "label": "Unique compound ID sample",
        "description": "Preview list of up to 50 distinct drug_id values linked to the protein.",
        "type": "multi_identifier",
        "separator": ";",
        "missing": "Empty means no binary compound relationship for this protein.",
        "role": "derived_summary",
        "source": "Sample from binary drug_id values.",
        "recommended": "Use only as a preview; binary table is authoritative.",
    },
    "has_matched_evidence": {
        "label": "Has matched evidence",
        "description": "Protein-level boolean indicating at least one matched structural/STITCH/experimental evidence signal.",
        "type": "boolean",
        "allowed": "0;1",
        "missing": "Should not be empty; 0 means no matched evidence captured.",
        "role": "derived_summary",
        "source": "Aggregated from binary evidence flags.",
        "recommended": "Use for quick protein-level evidence filtering.",
    },
    "reviewed": {
        "label": "UniProt reviewed status",
        "description": "UniProt reviewed/Swiss-Prot status when available.",
        "type": "category",
        "allowed": "reviewed",
        "missing": "Empty usually means unreviewed or not captured in enrichment.",
        "role": "annotation",
        "source": "UniProt enrichment.",
        "recommended": "Use to prioritize Swiss-Prot reviewed proteins.",
    },
    "go_cellular_component": {
        "label": "GO cellular component",
        "description": "GO cellular component annotations.",
        "type": "multi_category",
        "separator": ";",
        "missing": "Empty means no GO cellular component was captured.",
        "role": "annotation",
        "source": "UniProt/GO enrichment.",
        "recommended": "Use for localization context.",
    },
    "go_molecular_function": {
        "label": "GO molecular function",
        "description": "GO molecular function annotations.",
        "type": "multi_category",
        "separator": ";",
        "missing": "Empty means no GO molecular function was captured.",
        "role": "annotation",
        "source": "UniProt/GO enrichment.",
        "recommended": "Use for functional context.",
    },
    "pdb_structures": {
        "label": "PDB structure count",
        "description": "Count of linked PDB structures from UniProt enrichment.",
        "type": "integer",
        "missing": "Should not be empty; 0 means no linked PDB structure count captured.",
        "role": "annotation",
        "source": "UniProt xref_pdb enrichment.",
        "recommended": "Use to estimate existing structural coverage.",
    },
    "disease_association": {
        "label": "Disease association",
        "description": "UniProt disease association annotation text.",
        "type": "text",
        "missing": "Empty means no disease association was captured, not necessarily no disease relevance.",
        "role": "annotation",
        "source": "UniProt disease comments.",
        "recommended": "Use for biological context, not as exhaustive disease evidence.",
    },
    "transmembrane_count": {
        "label": "Transmembrane feature count",
        "description": "Count of UniProt TRANSMEM features.",
        "type": "integer",
        "missing": "Should not be empty; 0 means no transmembrane feature was counted.",
        "role": "annotation",
        "source": "UniProt TRANSMEM feature count.",
        "recommended": "Use to characterize membrane topology.",
    },
    "subcellular_location": {
        "label": "Subcellular location",
        "description": "UniProt subcellular location comment.",
        "type": "text",
        "missing": "Empty means no subcellular location comment was captured.",
        "role": "annotation",
        "source": "UniProt subcellular location comments.",
        "recommended": "Use for localization context.",
    },
    "protein_function": {
        "label": "Protein function",
        "description": "UniProt function comment.",
        "type": "text",
        "missing": "Empty means no function comment was captured.",
        "role": "annotation",
        "source": "UniProt function comments.",
        "recommended": "Use for protein-level functional interpretation.",
    },
    "source_row_count": {
        "label": "Relationship row count",
        "description": "Number of binary relationship rows involving this compound.",
        "type": "integer",
        "missing": "Should not be empty; 0 would mean no relationship rows.",
        "role": "derived_summary",
        "source": "Aggregated from binary table by drug_id.",
        "recommended": "Interpret as relationship/evidence row count, not unique target count.",
        "notes": "Consider renaming to relationship_row_count in a future cleanup.",
    },
    "unique_target_count": {
        "label": "Unique target count",
        "description": "Number of distinct target_uniprot_id values linked to this compound.",
        "type": "integer",
        "missing": "Should not be empty; 0 would mean no linked proteins.",
        "role": "derived_summary",
        "source": "Aggregated from binary table by drug_id.",
        "recommended": "Use for compound promiscuity/multi-target filtering.",
    },
    "activity_types": {
        "label": "Activity types",
        "description": "Semicolon-separated activity measurement types observed for the compound.",
        "type": "multi_category",
        "separator": ";",
        "missing": "Empty means no activity type was captured for this compound.",
        "role": "derived_summary",
        "source": "Aggregated from binary.activity_type.",
        "recommended": "Use to understand what quantitative evidence exists.",
    },
    "activity_value_uM_min": {
        "label": "Minimum activity value (uM)",
        "description": "Minimum numeric activity/affinity value across compound relationships.",
        "type": "float",
        "missing": "Empty means no numeric activity value was captured.",
        "role": "derived_summary",
        "source": "Aggregated from binary.activity_value_uM.",
        "recommended": "Use as a quick potency indicator with assay-context caution.",
    },
    "activity_value_uM_median": {
        "label": "Median activity value (uM)",
        "description": "Median numeric activity/affinity value across compound relationships.",
        "type": "float",
        "missing": "Empty means no numeric activity value was captured.",
        "role": "derived_summary",
        "source": "Aggregated from binary.activity_value_uM.",
        "recommended": "Use as a robust potency summary with assay-context caution.",
    },
    "molecular_formula": {
        "label": "Molecular formula",
        "description": "PubChem molecular formula.",
        "type": "text",
        "missing": "Empty means no PubChem property was retrieved, often for unmapped/biologic entries.",
        "role": "chemical_property",
        "source": "PubChem PUG REST property cache.",
        "recommended": "Use for chemistry filtering and identity checks.",
    },
    "molecular_weight": {
        "label": "Molecular weight",
        "description": "PubChem molecular weight in g/mol.",
        "type": "float",
        "missing": "Empty means no PubChem property was retrieved.",
        "role": "chemical_property",
        "source": "PubChem PUG REST property cache.",
        "recommended": "Use for drug-likeness and ligand-size filtering.",
    },
    "canonical_smiles": {
        "label": "Canonical SMILES",
        "description": "PubChem canonical SMILES string.",
        "type": "text",
        "missing": "Empty means no PubChem structure was retrieved.",
        "role": "chemical_structure",
        "source": "PubChem PUG REST property cache.",
        "recommended": "Use for cheminformatics where stereochemistry is not critical.",
    },
    "isomeric_smiles": {
        "label": "Isomeric SMILES",
        "description": "PubChem isomeric SMILES string with stereochemistry when available.",
        "type": "text",
        "missing": "Empty means no PubChem structure was retrieved.",
        "role": "chemical_structure",
        "source": "PubChem PUG REST property cache.",
        "recommended": "Prefer for cheminformatics when stereochemistry matters.",
    },
    "inchikey": {
        "label": "InChIKey",
        "description": "PubChem InChIKey.",
        "type": "identifier",
        "missing": "Empty means no PubChem structure mapping was available.",
        "role": "chemical_structure",
        "source": "PubChem PUG REST property cache.",
        "recommended": "Use for cross-database chemical identity matching.",
    },
    "iupac_name": {
        "label": "IUPAC name",
        "description": "PubChem IUPAC systematic name.",
        "type": "text",
        "missing": "Empty means no PubChem IUPAC name was retrieved.",
        "role": "chemical_property",
        "source": "PubChem PUG REST property cache.",
        "recommended": "Use for exact chemical naming; may be long.",
    },
    "xlogp": {
        "label": "XLogP",
        "description": "PubChem XLogP lipophilicity estimate.",
        "type": "float",
        "missing": "Empty means XLogP was not computed/retrieved.",
        "role": "chemical_property",
        "source": "PubChem PUG REST property cache.",
        "recommended": "Use for hydrophobicity/membrane-permeability heuristics.",
    },
    "tpsa": {
        "label": "TPSA",
        "description": "Topological polar surface area.",
        "type": "float",
        "missing": "Empty means property was not retrieved.",
        "role": "chemical_property",
        "source": "PubChem PUG REST property cache.",
        "recommended": "Use for permeability/drug-likeness filters.",
    },
    "hbond_donor_count": {
        "label": "Hydrogen bond donor count",
        "description": "PubChem hydrogen bond donor count.",
        "type": "integer",
        "missing": "Empty means property was not retrieved.",
        "role": "chemical_property",
        "source": "PubChem PUG REST property cache.",
        "recommended": "Use for drug-likeness filters.",
    },
    "hbond_acceptor_count": {
        "label": "Hydrogen bond acceptor count",
        "description": "PubChem hydrogen bond acceptor count.",
        "type": "integer",
        "missing": "Empty means property was not retrieved.",
        "role": "chemical_property",
        "source": "PubChem PUG REST property cache.",
        "recommended": "Use for drug-likeness filters.",
    },
    "rotatable_bond_count": {
        "label": "Rotatable bond count",
        "description": "PubChem rotatable bond count.",
        "type": "integer",
        "missing": "Empty means property was not retrieved.",
        "role": "chemical_property",
        "source": "PubChem PUG REST property cache.",
        "recommended": "Use for flexibility/drug-likeness filters.",
    },
    "complexity": {
        "label": "Molecular complexity",
        "description": "PubChem molecular complexity score.",
        "type": "float",
        "missing": "Empty means property was not retrieved.",
        "role": "chemical_property",
        "source": "PubChem PUG REST property cache.",
        "recommended": "Use as a rough structural complexity metric.",
    },
    "formal_charge": {
        "label": "Formal charge",
        "description": "PubChem formal charge.",
        "type": "integer",
        "missing": "Empty means property was not retrieved.",
        "role": "chemical_property",
        "source": "PubChem PUG REST property cache.",
        "recommended": "Use to identify charged compounds/ions.",
    },
    "ghs_hazard_classes": {
        "label": "GHS hazard classes",
        "description": "GHS hazard class text from PubChem/PUG-View when captured.",
        "type": "text",
        "missing": "`not_found_in_current_sources` means current sources did not provide this field; it does not prove no hazard.",
        "role": "low_coverage_safety",
        "source": "PubChem PUG-View safety/toxicity extraction.",
        "recommended": "Use cautiously as sparse safety annotation.",
    },
    "ghs_signal_words": {
        "label": "GHS signal words",
        "description": "GHS signal word such as Warning or Danger when captured.",
        "type": "category",
        "missing": "`not_found_in_current_sources` means current sources did not provide this field.",
        "role": "low_coverage_safety",
        "source": "PubChem PUG-View safety/toxicity extraction.",
        "recommended": "Currently very sparse; use cautiously.",
    },
    "toxicity_summary": {
        "label": "Toxicity summary",
        "description": "Human-readable toxicity summary text when captured.",
        "type": "text",
        "missing": "`not_found_in_current_sources` means current sources did not provide a toxicity summary; it does not prove no toxicity.",
        "role": "low_coverage_safety",
        "source": "PubChem PUG-View safety/toxicity extraction.",
        "recommended": "Use for qualitative review of known compounds only.",
    },
    "livertox": {
        "label": "LiverTox summary",
        "description": "LiverTox summary text when captured.",
        "type": "text",
        "missing": "`not_found_in_current_sources` means no LiverTox summary was captured.",
        "role": "low_coverage_safety",
        "source": "PubChem PUG-View/LiverTox extraction.",
        "recommended": "Use for liver-toxicity context in known drugs.",
    },
    "drug_classes": {
        "label": "Drug classes",
        "description": "Drug class/category text when captured.",
        "type": "text",
        "missing": "`not_found_in_current_sources` means no drug-class text was captured.",
        "role": "low_coverage_clinical",
        "source": "PubChem PUG-View clinical/pharmacology extraction.",
        "recommended": "Use as sparse clinical annotation.",
    },
    "pharmacodynamics": {
        "label": "Pharmacodynamics",
        "description": "Pharmacodynamic description when captured.",
        "type": "text",
        "missing": "`not_found_in_current_sources` means no pharmacodynamic text was captured.",
        "role": "low_coverage_clinical",
        "source": "PubChem PUG-View clinical/pharmacology extraction.",
        "recommended": "Use for qualitative pharmacology review.",
    },
    "drug_indication": {
        "label": "Drug indication",
        "description": "Indication/use description when captured.",
        "type": "text",
        "missing": "`not_found_in_current_sources` means no indication text was captured; it does not prove no medical use.",
        "role": "low_coverage_clinical",
        "source": "PubChem PUG-View clinical/pharmacology extraction.",
        "recommended": "Use as sparse clinical annotation.",
    },
    "pocket_batch": {
        "label": "Pocket prediction batch",
        "description": "Batch label for AlphaFold pocket prediction input.",
        "type": "category",
        "allowed": "1;2",
        "missing": "Should not be empty.",
        "role": "audit",
        "source": "AlphaFold pocket prediction batch metadata.",
        "recommended": "Use only for provenance/QC.",
    },
    "pocket_id": {
        "label": "Pocket ID",
        "description": "Unique pocket identifier, usually UniProt plus pocket number.",
        "type": "identifier",
        "missing": "Required pocket key; should not be empty.",
        "role": "core_identifier",
        "source": "AlphaFold pocket prediction pipeline.",
        "recommended": "Use as primary key for pocket instances.",
    },
    "center_x": {
        "label": "Pocket center X",
        "description": "X coordinate of predicted pocket center.",
        "type": "float",
        "missing": "Should not be empty.",
        "role": "pocket_geometry",
        "source": "AlphaFold pocket prediction pipeline.",
        "recommended": "Use for spatial pocket analysis.",
    },
    "center_y": {
        "label": "Pocket center Y",
        "description": "Y coordinate of predicted pocket center.",
        "type": "float",
        "missing": "Should not be empty.",
        "role": "pocket_geometry",
        "source": "AlphaFold pocket prediction pipeline.",
        "recommended": "Use for spatial pocket analysis.",
    },
    "center_z": {
        "label": "Pocket center Z",
        "description": "Z coordinate of predicted pocket center.",
        "type": "float",
        "missing": "Should not be empty.",
        "role": "pocket_geometry",
        "source": "AlphaFold pocket prediction pipeline.",
        "recommended": "Use for spatial pocket analysis.",
    },
    "n_residues": {
        "label": "Pocket residue count",
        "description": "Number of residues assigned to the predicted pocket.",
        "type": "integer",
        "missing": "Should not be empty.",
        "role": "pocket_geometry",
        "source": "AlphaFold pocket prediction pipeline.",
        "recommended": "Use to filter pocket size.",
    },
    "score": {
        "label": "Pocket score",
        "description": "Heuristic pocket score from the AlphaFold pocket prediction pipeline.",
        "type": "float",
        "missing": "Should not be empty.",
        "role": "pocket_geometry",
        "source": "AlphaFold pocket prediction pipeline.",
        "recommended": "Use to rank predicted pockets within a protein, with caution.",
    },
    "pocket_residues": {
        "label": "Pocket residues",
        "description": "Semicolon-separated residue list assigned to the predicted pocket.",
        "type": "multi_category",
        "separator": ";",
        "missing": "Should not be empty.",
        "role": "pocket_geometry",
        "source": "AlphaFold pocket prediction pipeline.",
        "recommended": "Use for residue-level pocket comparison.",
    },
    "source_file": {
        "label": "Pocket source file",
        "description": "Original pocket prediction CSV source file.",
        "type": "text",
        "missing": "Should not be empty.",
        "role": "audit",
        "source": "Pocket prediction pipeline output filename.",
        "recommended": "Use for provenance/debugging.",
    },
}


TABLE_SPECIFIC: dict[tuple[str, str], dict[str, str]] = {
    ("2_Protein_Database", "target_uniprot_id"): {
        "primary": "yes",
        "foreign": "no",
        "references": "",
        "notes": "One row per UniProt protein.",
    },
    ("3_Small_Molecules", "drug_id"): {
        "primary": "yes",
        "foreign": "no",
        "references": "",
        "notes": "One row per compound/drug identifier.",
    },
    ("4_Pocket_Instances", "pocket_id"): {
        "primary": "yes",
        "foreign": "no",
        "references": "",
        "notes": "One row per AlphaFold predicted pocket instance.",
    },
}


def table_context(sheet: str, col: str) -> dict[str, str]:
    meta: dict[str, str] = {}
    if col == "target_uniprot_id" and sheet in {"1_Binary_Relationships", "4_Pocket_Instances"}:
        meta.update({"primary": "no", "foreign": "yes", "references": "2_Protein_Database.target_uniprot_id"})
    if col == "drug_id" and sheet == "1_Binary_Relationships":
        meta.update({"primary": "no", "foreign": "yes", "references": "3_Small_Molecules.drug_id"})
    if sheet == "1_Binary_Relationships":
        meta.setdefault("primary", "no")
        meta.setdefault("foreign", "no")
        meta.setdefault("references", "")
        if col in {"target_uniprot_id", "drug_id", "source_database", "assay_or_mechanism", "activity_type", "pdb_biolip_matched_sites", "pdb_scpdb_matched_sites", "pdbbind_matched_sites"}:
            meta.setdefault("notes", "Binary table rows are relationship/context rows; uniqueness is contextual, not a single-column primary key.")
    meta.update(TABLE_SPECIFIC.get((sheet, col), {}))
    return meta


def infer_type(col: str, values: list[str]) -> str:
    if col.startswith("has_") or col == "has_matched_evidence":
        return "boolean"
    if col.endswith("_count") or col in {"source_row_count", "unique_target_count", "unique_drug_count", "n_residues", "pdb_structures", "transmembrane_count", "n_pockets_alphafold"}:
        return "integer"
    if col.startswith("center_") or col in {"score", "activity_value_uM", "activity_value_uM_min", "activity_value_uM_median", "molecular_weight", "xlogp", "tpsa", "complexity"}:
        return "float"
    if any(";" in v for v in values if v):
        return "multi_category"
    return "text"


def summarize_column(values: list[str]) -> dict[str, Any]:
    total = len(values)
    not_found = sum(1 for v in values if v == NOT_FOUND)
    empty = sum(1 for v in values if v == "")
    real_values = [v for v in values if v not in {"", NOT_FOUND}]
    real_count = len(real_values)
    counter = Counter(real_values)
    top_values = "; ".join(f"{k} ({v})" for k, v in counter.most_common(8))
    examples = []
    seen = set()
    for value in real_values:
        if value not in seen:
            examples.append(value)
            seen.add(value)
        if len(examples) >= 5:
            break
    return {
        "total": total,
        "not_found": not_found,
        "empty": empty,
        "real_count": real_count,
        "coverage": round(real_count / total * 100, 2) if total else 0,
        "unique": len(counter),
        "top_values": top_values,
        "examples": "; ".join(examples),
    }


def default_meta(col: str, values: list[str]) -> dict[str, str]:
    return {
        "label": col.replace("_", " "),
        "description": f"Field `{col}` in the normalized dataset.",
        "type": infer_type(col, values),
        "allowed": "",
        "separator": ";" if any(";" in v for v in values if v) else "",
        "missing": "Empty means no value was captured in the current sources.",
        "role": "annotation",
        "source": "Derived from the normalized table generation pipeline.",
        "recommended": "Review field-specific documentation before analytical use.",
        "notes": "",
    }


def allowed_values_sheet_rows() -> list[dict[str, str]]:
    rows = []
    for col, values in ALLOWED_VALUE_NOTES.items():
        for value, meaning in values.items():
            rows.append({"column_name": col, "allowed_value": value, "meaning": meaning})
    return rows


def build_rows() -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for sheet, file_name in TABLES:
        path = BASE / file_name
        with path.open("r", encoding="utf-8", newline="") as f:
            reader = csv.DictReader(f, delimiter="\t")
            data = list(reader)
            header = reader.fieldnames or []
        by_col = {col: [row.get(col, "") for row in data] for col in header}
        for idx, col in enumerate(header, 1):
            values = by_col[col]
            meta = default_meta(col, values)
            meta.update(BASE_META.get(col, {}))
            meta.update(table_context(sheet, col))
            summary = summarize_column(values)
            rows.append(
                {
                    "sheet_name": sheet,
                    "table_file": file_name,
                    "column_order": idx,
                    "column_name": col,
                    "column_label": meta.get("label", ""),
                    "description": meta.get("description", ""),
                    "data_type": meta.get("type", ""),
                    "is_primary_key": meta.get("primary", "no"),
                    "is_foreign_key": meta.get("foreign", "no"),
                    "references": meta.get("references", ""),
                    "allowed_values": meta.get("allowed", ""),
                    "separator": meta.get("separator", ""),
                    "missing_value_meaning": meta.get("missing", ""),
                    "field_role": meta.get("role", ""),
                    "source_or_rule": meta.get("source", ""),
                    "total_rows": summary["total"],
                    "real_value_count": summary["real_count"],
                    "not_found_count": summary["not_found"],
                    "empty_count": summary["empty"],
                    "coverage_percent": summary["coverage"],
                    "observed_unique_count": summary["unique"],
                    "top_values": summary["top_values"],
                    "example_values": summary["examples"],
                    "notes": meta.get("notes", ""),
                    "recommended_use": meta.get("recommended", ""),
                }
            )
    return rows


def write_tsv(rows: list[dict[str, Any]]) -> None:
    with OUT_TSV.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=DICT_COLUMNS, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(rows: list[dict[str, Any]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data_Dictionary"
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    ws.append(DICT_COLUMNS)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        ws.append([row.get(col, "") for col in DICT_COLUMNS])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(DICT_COLUMNS))}{len(rows)+1}"
    widths = {
        "A": 24,
        "B": 34,
        "C": 12,
        "D": 32,
        "E": 28,
        "F": 58,
        "G": 16,
        "H": 14,
        "I": 14,
        "J": 38,
        "K": 48,
        "L": 12,
        "M": 52,
        "N": 22,
        "O": 56,
        "P": 12,
        "Q": 16,
        "R": 16,
        "S": 12,
        "T": 14,
        "U": 18,
        "V": 60,
        "W": 60,
        "X": 52,
        "Y": 54,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws2 = wb.create_sheet("Allowed_Values")
    allowed_cols = ["column_name", "allowed_value", "meaning"]
    ws2.append(allowed_cols)
    for cell in ws2[1]:
        cell.fill = header_fill
        cell.font = header_font
    for row in allowed_values_sheet_rows():
        ws2.append([row[col] for col in allowed_cols])
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:C{ws2.max_row}"
    ws2.column_dimensions["A"].width = 38
    ws2.column_dimensions["B"].width = 42
    ws2.column_dimensions["C"].width = 90
    for row in ws2.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws3 = wb.create_sheet("Field_Roles")
    role_counts = Counter(row["field_role"] for row in rows)
    ws3.append(["field_role", "field_count", "meaning"])
    role_meanings = {
        "core_identifier": "Primary/join identifier fields.",
        "core_annotation": "Human-readable names and core labels.",
        "cross_reference": "External database identifiers.",
        "evidence_source": "Source/provenance fields for evidence.",
        "relationship_evidence": "Relationship-level assay/mechanism evidence.",
        "activity_evidence": "Quantitative activity/affinity evidence.",
        "binding_site_evidence": "Structure/site evidence fields.",
        "network_evidence": "Network or interaction evidence fields.",
        "binding_site_annotation": "Protein-level experimental site annotations.",
        "literature_reference": "Literature traceability fields.",
        "clinical_status": "Clinical or approval status fields.",
        "derived_classification": "Derived class/status fields.",
        "derived_interpretation": "Derived ligand interpretation fields.",
        "derived_summary": "Entity-level summaries derived from row-level tables.",
        "annotation": "External annotation/enrichment fields.",
        "chemical_property": "PubChem chemical property fields.",
        "chemical_structure": "Chemical structure representation fields.",
        "low_coverage_safety": "Sparse safety/toxicity annotation fields.",
        "low_coverage_clinical": "Sparse clinical/pharmacology text fields.",
        "pocket_geometry": "AlphaFold predicted pocket geometry fields.",
        "reserved": "Reserved or currently unpopulated fields.",
        "audit": "Provenance/debugging/audit fields.",
    }
    for role, count in sorted(role_counts.items()):
        ws3.append([role, count, role_meanings.get(role, "")])
    ws3.freeze_panes = "A2"
    ws3.auto_filter.ref = f"A1:C{ws3.max_row}"
    ws3.column_dimensions["A"].width = 30
    ws3.column_dimensions["B"].width = 14
    ws3.column_dimensions["C"].width = 90
    for row in ws3.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in ws3[1]:
        cell.fill = header_fill
        cell.font = header_font

    wb.save(OUT_XLSX)


def main() -> None:
    csv.field_size_limit(100_000_000)
    rows = build_rows()
    write_tsv(rows)
    write_xlsx(rows)
    print(json.dumps({"rows": len(rows), "tsv": str(OUT_TSV), "xlsx": str(OUT_XLSX)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
