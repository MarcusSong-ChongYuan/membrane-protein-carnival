import csv
import hashlib
import json
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(r"C:\github-repos\upload\normalized_tables")
BINARY = BASE / "drug_protein_binary_relationships.tsv"
PROTEIN = BASE / "protein_database.tsv"
SMALL = BASE / "small_molecule_database.tsv"
REVIEW = BASE / "protein_structure_ligand_context_review.tsv"
OUT_TSV = BASE / "protein_structure_ligand_context_binary_format_review.tsv"
OUT_XLSX = BASE / "protein_structure_ligand_context_binary_format_review.xlsx"
MANIFEST = BASE / "MANIFEST.json"

csv.field_size_limit(100_000_000)


def load_tsv(path):
    with path.open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f, delimiter="\t"))


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def site_value(row):
    raw = row.get("raw_entry", "")
    cid = row.get("ligand_pubchem_cid", "")
    prefix = f"CID{cid}|" if cid else "CID_UNMAPPED|"
    return prefix + raw


def mechanism(row):
    parts = [
        "structure_ligand_context",
        f"category={row.get('ligand_category', '')}",
        f"cid_status={row.get('cid_mapping_status', '')}",
        f"keep_small={row.get('keep_for_small_molecule_binding_site', '')}",
        f"keep_context={row.get('keep_for_binding_site_context', '')}",
    ]
    if row.get("exclusion_reason"):
        parts.append(f"exclusion={row['exclusion_reason']}")
    if row.get("structure_ligand_id"):
        parts.append(f"context_id={row['structure_ligand_id']}")
    return "; ".join(parts)


def main():
    binary_rows = load_tsv(BINARY)
    binary_fields = list(binary_rows[0].keys())
    protein_rows = load_tsv(PROTEIN)
    small_rows = load_tsv(SMALL)
    review_rows = load_tsv(REVIEW)

    protein_by_uid = {row["target_uniprot_id"]: row for row in protein_rows}
    small_by_cid = {row["drug_id"]: row for row in small_rows}

    output_rows = []
    for row in review_rows:
        uid = row["target_uniprot_id"]
        protein = protein_by_uid.get(uid, {})
        cid = row.get("ligand_pubchem_cid", "")
        small = small_by_cid.get(cid, {})
        source = row.get("structure_source", "")

        out = {field: "" for field in binary_fields}
        out["target_uniprot_id"] = uid
        out["approved_symbol"] = row.get("approved_symbol") or protein.get("approved_symbol", "")
        out["ncbi_gene_id"] = protein.get("ncbi_gene_id", "")
        out["protein_name"] = protein.get("protein_name", "")
        out["drug_id"] = cid
        out["compound_id_type"] = "PubChem CID" if cid else "unmapped_structure_ligand"
        out["drug_name"] = (
            small.get("drug_name")
            or row.get("ligand_name_raw")
            or row.get("raw_ligand_label", "")
        )
        out["drug_name_type"] = small.get("drug_name_type") or row.get("ligand_category", "")
        out["compound_source"] = "structure_context"
        out["source_database"] = source
        out["assay_or_mechanism"] = mechanism(row)
        out["activity_type"] = row.get("affinity_type", "")
        out["activity_value_uM"] = ""
        out["clinical_or_approval_status"] = ""

        if row.get("affinity_value"):
            # Keep original value/unit in the mechanism field unless it is already uM;
            # this avoids silently converting nM and other units in a review file.
            unit = row.get("affinity_unit", "")
            if unit.lower() == "um":
                out["activity_value_uM"] = row["affinity_value"]
            else:
                out["assay_or_mechanism"] += f"; affinity_original={row['affinity_value']}{unit}"

        if source == "BioLiP":
            out["has_pdb_biolip_matched"] = "1"
            out["pdb_biolip_matched_sites"] = site_value(row)
        else:
            out["has_pdb_biolip_matched"] = "0"

        if source == "sc-PDB":
            out["has_pdb_scpdb_matched"] = "1"
            out["pdb_scpdb_matched_sites"] = site_value(row)
        else:
            out["has_pdb_scpdb_matched"] = "0"

        if source == "PDBbind":
            out["has_pdbbind_matched"] = "1"
            out["pdbbind_matched_sites"] = site_value(row)
        else:
            out["has_pdbbind_matched"] = "0"

        out["has_stitch_matched"] = "0"
        out["has_exp_binding_site"] = "0"
        output_rows.append(out)

    with OUT_TSV.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, delimiter="\t", fieldnames=binary_fields, lineterminator="\n")
        writer.writeheader()
        writer.writerows(output_rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "Binary_Format_Review"
    ws.append(binary_fields)
    for row in output_rows:
        ws.append([row[field] for field in binary_fields])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(binary_fields))}{len(output_rows) + 1}"
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col_idx, field in enumerate(binary_fields, 1):
        max_len = len(field)
        for row_idx in range(2, min(len(output_rows) + 2, 250)):
            max_len = max(max_len, len(str(ws.cell(row_idx, col_idx).value or "")))
        ws.column_dimensions[get_column_letter(col_idx)].width = max(10, min(max_len + 2, 45))
    wb.save(OUT_XLSX)

    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    files = manifest.get("files", [])
    by_name = {Path(item["path"]).name: item for item in files}
    for path in [OUT_TSV, OUT_XLSX]:
        item = by_name.get(path.name)
        if item is None:
            item = {"path": f"upload/normalized_tables/{path.name}"}
            files.append(item)
        item["sha256"] = sha256(path)
        item["bytes"] = path.stat().st_size
    manifest["files"] = files
    MANIFEST.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    print(json.dumps({
        "rows": len(output_rows),
        "cols": len(binary_fields),
        "mapped_cid_rows": sum(1 for row in output_rows if row["drug_id"]),
        "unmapped_rows": sum(1 for row in output_rows if not row["drug_id"]),
        "output_tsv": str(OUT_TSV),
        "output_xlsx": str(OUT_XLSX),
    }, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
