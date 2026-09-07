from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

import fitz
from docx import Document
from docx.oxml.ns import qn


ROOT = Path(r"C:\Users\Administrator\HuMemLigDB_manuscript")
DOCX = ROOT / "HuMemLigDB_NAR_manuscript_draft_v0.1.docx"
PDF = ROOT / "rendered_v0_1" / "HuMemLigDB_NAR_manuscript_draft_v0.1_qa2.pdf"
REPORT = ROOT / "HuMemLigDB_NAR_manuscript_draft_v0.1_QA.json"


def compact(text: str, limit: int = 150) -> str:
    text = re.sub(r"\s+", " ", text).strip()
    return text[:limit]


def pdf_qa() -> dict:
    pdf = fitz.open(PDF)
    pages = []
    global_errors = []
    for index, page in enumerate(pdf, start=1):
        rect = page.rect
        raw = page.get_text("dict")
        spans = []
        blocks = []
        for block in raw.get("blocks", []):
            if block.get("type") != 0:
                continue
            block_text = []
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    text = span.get("text", "")
                    if text.strip():
                        spans.append(span)
                        block_text.append(text)
            if block_text:
                blocks.append(
                    {
                        "bbox": [round(v, 2) for v in block.get("bbox", [0, 0, 0, 0])],
                        "text": compact(" ".join(block_text)),
                    }
                )
        words = page.get_text("words")
        outside = []
        for word in words:
            x0, y0, x1, y1, txt = word[:5]
            if x0 < -0.5 or y0 < -0.5 or x1 > rect.width + 0.5 or y1 > rect.height + 0.5:
                outside.append([round(x0, 1), round(y0, 1), round(x1, 1), round(y1, 1), txt])
        sizes = [round(float(s.get("size", 0)), 2) for s in spans if float(s.get("size", 0)) > 0]
        fonts = sorted({s.get("font", "") for s in spans if s.get("font")})
        text = page.get_text("text")
        lines = [compact(line, 180) for line in text.splitlines() if line.strip()]
        page_errors = []
        if not words:
            page_errors.append("blank_page")
        if outside:
            page_errors.append("text_outside_page")
        if sizes and min(sizes) < 7.5:
            page_errors.append("font_below_7_5_pt")
        # A heading should not be stranded at the bottom with no following prose.
        bottom_blocks = [b for b in blocks if b["bbox"][1] > rect.height - 90]
        for block in bottom_blocks:
            if len(block["text"]) < 90 and re.match(r"^(Figure|Table|[A-Z][A-Za-z -]{2,})", block["text"]):
                page_errors.append("possible_orphan_heading_near_bottom")
                break
        pages.append(
            {
                "page": index,
                "size_points": [rect.width, rect.height],
                "word_count": len(words),
                "block_count": len(blocks),
                "min_font_pt": min(sizes) if sizes else None,
                "max_font_pt": max(sizes) if sizes else None,
                "fonts": fonts,
                "first_lines": lines[:4],
                "last_lines": lines[-4:],
                "outside_words": outside[:10],
                "errors": sorted(set(page_errors)),
            }
        )
        global_errors.extend(f"page_{index}:{x}" for x in sorted(set(page_errors)))
    full_text = "\n".join(page.get_text("text") for page in pdf)
    prohibited = [
        token
        for token in ("294,?", "[This source-count", "**", "```", "PLACEHOLDER")
        if token in full_text
    ]
    required = [
        "Abstract",
        "Introduction",
        "Results",
        "Discussion",
        "Materials and methods",
        "Data availability",
        "1,305,791",
        "7,904",
        "3,003,306",
    ]
    missing = [token for token in required if token not in full_text]
    return {
        "page_count": len(pdf),
        "prohibited_tokens_found": prohibited,
        "required_tokens_missing": missing,
        "global_errors": global_errors,
        "pages": pages,
    }


def docx_qa() -> dict:
    doc = Document(DOCX)
    section = doc.sections[0]
    text = "\n".join(p.text for p in doc.paragraphs)
    tables = []
    errors = []
    for idx, table in enumerate(doc.tables, start=1):
        grid = table._tbl.tblGrid
        widths = [int(x.get(qn("w:w"))) for x in grid.findall(qn("w:gridCol"))]
        tbl_w = table._tbl.tblPr.find(qn("w:tblW"))
        tbl_ind = table._tbl.tblPr.find(qn("w:tblInd"))
        width = int(tbl_w.get(qn("w:w"))) if tbl_w is not None else None
        indent = int(tbl_ind.get(qn("w:w"))) if tbl_ind is not None else None
        if sum(widths) != 9360 or width != 9360:
            errors.append(f"table_{idx}:geometry")
        if indent != 120:
            errors.append(f"table_{idx}:indent")
        tables.append(
            {
                "table": idx,
                "rows": len(table.rows),
                "cols": len(table.columns),
                "grid_widths": widths,
                "grid_sum": sum(widths),
                "tblW": width,
                "tblInd": indent,
            }
        )
    if "294,?" in text or "[This source-count" in text:
        errors.append("unresolved_source_count_placeholder")
    if "1,305,791" not in text:
        errors.append("core_pair_count_missing")
    with zipfile.ZipFile(DOCX) as archive:
        bad = archive.testzip()
        if bad:
            errors.append(f"zip_crc:{bad}")
    return {
        "page_size_inches": [section.page_width.inches, section.page_height.inches],
        "margins_inches": [
            section.top_margin.inches,
            section.right_margin.inches,
            section.bottom_margin.inches,
            section.left_margin.inches,
        ],
        "paragraphs": len(doc.paragraphs),
        "tables": tables,
        "errors": errors,
    }


def main() -> None:
    result = {
        "docx": str(DOCX),
        "pdf": str(PDF),
        "docx_qa": docx_qa(),
        "pdf_qa": pdf_qa(),
    }
    result["status"] = (
        "PASS"
        if not result["docx_qa"]["errors"]
        and not result["pdf_qa"]["prohibited_tokens_found"]
        and not result["pdf_qa"]["required_tokens_missing"]
        and not any("text_outside_page" in x or "blank_page" in x for x in result["pdf_qa"]["global_errors"])
        else "REVIEW"
    )
    REPORT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(REPORT)
    print(result["status"])
    print("\n".join(result["pdf_qa"]["global_errors"]))


if __name__ == "__main__":
    main()
