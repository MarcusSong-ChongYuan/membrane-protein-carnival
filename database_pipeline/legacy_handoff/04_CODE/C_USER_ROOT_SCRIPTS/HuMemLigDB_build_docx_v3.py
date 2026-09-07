from __future__ import annotations

from docx import Document
from docx.shared import Pt

import HuMemLigDB_build_docx as base
import HuMemLigDB_build_docx_v2  # applies the verified content normalization


def is_page_break_only(paragraph) -> bool:
    return not paragraph.text.strip() and bool(paragraph._p.xpath(".//w:br[@w:type='page']"))


def remove_duplicate_section_breaks(doc: Document) -> None:
    for heading_text in ("Tables", "Figure legends"):
        paragraphs = list(doc.paragraphs)
        for idx, paragraph in enumerate(paragraphs[:-1]):
            if paragraph.text.strip() == heading_text and is_page_break_only(paragraphs[idx + 1]):
                node = paragraphs[idx + 1]._p
                node.getparent().remove(node)
                break


def compact_reference_completion(doc: Document) -> None:
    prefix = "The final bibliography should include"
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.10
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
            break


def main() -> None:
    base.main()
    doc = Document(base.OUTPUT)
    remove_duplicate_section_breaks(doc)
    compact_reference_completion(doc)
    base.audit_document(doc)
    doc.save(base.OUTPUT)
    print(base.OUTPUT)


if __name__ == "__main__":
    main()
